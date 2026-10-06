import json
import html
import base64
import difflib
import hashlib
import hmac
import ipaddress
import os
import re
import shutil
import subprocess
import sys
import tempfile
import tkinter as tk
from datetime import datetime, timedelta, timezone
from pathlib import Path
from tkinter import filedialog, messagebox, simpledialog, ttk
from urllib.parse import urlparse
from PIL import Image, ImageTk
from docx import Document

WINDOWS_HOSTS_PATH = r"C:\Windows\System32\drivers\etc\hosts"
APP_ROOT = Path(__file__).resolve().parent
RESOURCE_ROOT = Path(getattr(sys, "_MEIPASS", APP_ROOT))
LOCAL_APP_DATA = Path(os.environ.get("LOCALAPPDATA", str(Path.home() / "AppData" / "Local")))
AGREEMENT_FILE = LOCAL_APP_DATA / "Avantis FireWall" / "agreement.json"
AGREEMENT_VERSION = "1.3"
APP_ICON_FILE = RESOURCE_ROOT / "Images" / "avantis-app.ico"
if getattr(sys, "frozen", False) and os.name == "nt":
    PROGRAM_DATA = Path(os.environ.get("PROGRAMDATA", r"C:\ProgramData"))
    LEGACY_DATA_ROOT = PROGRAM_DATA / "Avantis FireWall"
    DATA_ROOT = LOCAL_APP_DATA / "Avantis FireWall"
else:
    LEGACY_DATA_ROOT = None
    DATA_ROOT = APP_ROOT
HOSTS_BACKUP_PATH = DATA_ROOT / "hosts.avantis.backup"
RULES_FILE = DATA_ROOT / "rules.json"
ACTIVITY_FILE = DATA_ROOT / "activity.json"
POLICY_HISTORY_FILE = DATA_ROOT / "policy-history.json"
ADMIN_AUDIT_FILE = DATA_ROOT / "admin-audit.json"
LOGO_FILE = RESOURCE_ROOT / "Images" / "Avantis-logo-prl.png"
MIGRATABLE_DATA_FILES = (
    "rules.json",
    "activity.json",
    "policy-history.json",
    "admin-audit.json",
    "hosts.avantis.backup",
)
DNS_BLOCKLIST_FILENAMES = (
    "avantis-dns-blocklist",
    "avantis-dns-blocklist.txt",
    "avantis dns blocklist.txt",
)
PROFILE_LABELS = {"default": "Admin", "child_protection": "Safe Mode"}
PROFILE_VALUES = {label: value for value, label in PROFILE_LABELS.items()}
PROFILE_PASSWORD_HASH_PREFIX = "pbkdf2_sha256$"
PROFILE_PASSWORD_HASH_ITERATIONS = 600_000
DOMAIN_PATTERN = re.compile(r"(?i)(?:https?://)?(?:www\.)?[a-z0-9](?:[a-z0-9.-]*[a-z0-9])?\.[a-z]{2,}")
VALID_DOMAIN_PATTERN = re.compile(r"(?i)(?:[a-z0-9](?:[a-z0-9-]{0,61}[a-z0-9])?\.)+[a-z](?:[a-z0-9-]{0,61}[a-z0-9])?")

DEFAULT_CATEGORIES = {
    "adult": [
        "adult", "porn", "porno", "pornhub", "sex", "xxx", "xvideos",
        "nude", "nudity", "strip", "erotic", "playboy", "asian porn",
        "pornography", "explicit", "xxx video", "sex cam", "cam girl",
        "escort", "hookup"
    ],
    "gambling": [
        "casino", "gambling", "betting", "roulette", "poker", "slot",
        "jackpot", "lottery", "sportsbook", "baccarat", "sports betting",
        "live betting", "virtual betting", "betting tips", "casino online",
        "aviator", "crash game", "bet slip", "odds", "bookmaker", "bookie",
        "slot games"
    ],
    "malware": [
        "crack", "keygen", "warez", "nulled", "r57", "download free",
        "torrent", "serial key", "activator", "pirated", "cheat", "mod apk"
    ],
    "scam": [
        "free money", "bitcoin generator", "lottery winner", "fake rewards",
        "clickbank scam", "crypto scam", "giveaway", "claim reward",
        "investment scam", "double your money", "urgent payment", "verify account"
    ],
}

def _parse_profile_password_hash(value):
    try:
        scheme, iterations_text, salt_text, digest_text = value.split("$")
        iterations = int(iterations_text)
        salt = base64.b64decode(salt_text + "=" * (-len(salt_text) % 4), altchars=b"-_", validate=True)
        digest = base64.b64decode(digest_text + "=" * (-len(digest_text) % 4), altchars=b"-_", validate=True)
    except (AttributeError, TypeError, ValueError):
        return None
    if (
        scheme != "pbkdf2_sha256"
        or not 100_000 <= iterations <= 2_000_000
        or len(salt) != 16
        or len(digest) != hashlib.sha256().digest_size
    ):
        return None
    return iterations, salt, digest


def hash_profile_password(password):
    password = str(password or "")
    if not password:
        return ""
    if password.startswith(PROFILE_PASSWORD_HASH_PREFIX):
        if _parse_profile_password_hash(password) is None:
            raise ValueError("The saved Safe Mode password hash is invalid.")
        return password

    salt = os.urandom(16)
    digest = hashlib.pbkdf2_hmac(
        "sha256",
        password.encode("utf-8"),
        salt,
        PROFILE_PASSWORD_HASH_ITERATIONS,
    )
    salt_text = base64.urlsafe_b64encode(salt).decode("ascii").rstrip("=")
    digest_text = base64.urlsafe_b64encode(digest).decode("ascii").rstrip("=")
    return f"{PROFILE_PASSWORD_HASH_PREFIX}{PROFILE_PASSWORD_HASH_ITERATIONS}${salt_text}${digest_text}"


def verify_profile_password(password, stored_value):
    password = str(password or "")
    stored_value = str(stored_value or "")
    if stored_value.startswith(PROFILE_PASSWORD_HASH_PREFIX):
        parsed = _parse_profile_password_hash(stored_value)
        if parsed is None:
            return False
        iterations, salt, expected_digest = parsed
        actual_digest = hashlib.pbkdf2_hmac("sha256", password.encode("utf-8"), salt, iterations)
        return hmac.compare_digest(actual_digest, expected_digest)
    return hmac.compare_digest(password, stored_value)


DEFAULT_BLOCKLIST = [
    "pornhub.com",
    "xvideos.com",
    "xnxx.com",
    "bet365.com",
    "casino.com",
    "888casino.com",
    "pokerstars.com",
    "betway.co.zw",
    "hollywoodbets.co.zw",
    "sportingbet.co.zw",
    "zimbet.co.zw",
    "1xbet.co.zw",
    "parimatch.co.zw",
    "stake.com",
    "1xbet.com",
    "betway.com",
    "betano.com",
    "betfair.com",
    "unibet.com",
    "bwin.com",
    "ladbrokes.com",
    "williamhill.com",
    "paddypower.com",
    "betvictor.com",
    "22bet.com",
    "melbet.com",
    "mostbet.com",
    "roobet.com",
    "rollbit.com",
    "duckbet.com",
    "redtube.com",
    "youporn.com",
    "xhamster.com",
    "xnxx.tv",
    "spankbang.com",
    "chaturbate.com",
    "cam4.com",
    "onlyfans.com",
]

DEFAULT_SUBDOMAIN_PREFIXES = [
    "www", "info", "m", "mobile", "login", "account", "sports",
    "casino", "bet", "api", "app", "play",
]

CATEGORY_DESCRIPTIONS = {
    "adult": "Adult or sexually explicit sites and terms.",
    "gambling": "Casino, betting, lottery, and wagering sites and terms.",
    "malware": "Terms commonly associated with pirated or unsafe downloads.",
    "scam": "Fraud, fake rewards, and suspicious money-request terms.",
}

RETENTION_OPTIONS = {
    "1 hour": 1,
    "24 hours": 24,
    "7 days": 24 * 7,
}

SHORTENER_DOMAINS = {
    "bit.ly", "tinyurl.com", "t.co", "is.gd", "ow.ly", "buff.ly", "cutt.ly",
}
SUSPICIOUS_TLDS = {"click", "download", "gq", "loan", "ml", "tk", "top", "work", "zip"}
PUBLIC_LIST_LIMIT = 5000
PROTECTED_HOSTS = {
    "localhost", "127.0.0.1", "microsoft.com", "google.com", "bing.com",
    "duckduckgo.com", "yahoo.com", "github.com", "openai.com", "copilot.microsoft.com",
}
PUBLIC_LISTS = {
    "URLhaus malware hostfile": {
        "url": "https://urlhaus.abuse.ch/downloads/hostfile/",
        "description": "Abuse.ch hostnames associated with malware distribution.",
    },
    "StevenBlack hosts list": {
        "url": "https://raw.githubusercontent.com/StevenBlack/hosts/master/hosts",
        "description": "Community-maintained hosts list containing ads and known harmful domains.",
    },
}
HOSTS_DOMAIN_LIMIT = 1000
PROTECTED_DOMAINS = {
    "bing.com", "google.com", "yahoo.com", "duckduckgo.com",
    "microsoft.com", "windows.com", "live.com", "openai.com", "chatgpt.com",
}


def default_rules():
    return {
        "blocked_domains": DEFAULT_BLOCKLIST[:],
        "categories": {category: keywords[:] for category, keywords in DEFAULT_CATEGORIES.items()},
        "subdomain_prefixes": DEFAULT_SUBDOMAIN_PREFIXES[:],
        "protection_profile": "default",
        "profile_password": "",
    }


def initialize_data_root():
    DATA_ROOT.mkdir(parents=True, exist_ok=True)
    if LEGACY_DATA_ROOT is None or not LEGACY_DATA_ROOT.is_dir():
        return

    for name in MIGRATABLE_DATA_FILES:
        source = LEGACY_DATA_ROOT / name
        destination = DATA_ROOT / name
        if source.is_file() and not destination.exists():
            shutil.copy2(source, destination)


def load_config():
    if not RULES_FILE.exists():
        return default_rules()

    try:
        with RULES_FILE.open("r", encoding="utf-8") as f:
            data = json.load(f)

        # Keep support for the original rules.json format: ["example.com"].
        if isinstance(data, list):
            config = default_rules()
            config["blocked_domains"] = data
            return config

        if isinstance(data, dict):
            config = default_rules()
            if isinstance(data.get("blocked_domains"), list):
                config["blocked_domains"] = data["blocked_domains"]
            if isinstance(data.get("categories"), dict):
                config["categories"] = {
                    str(category): [str(keyword) for keyword in keywords if str(keyword).strip()]
                    for category, keywords in data["categories"].items()
                    if isinstance(keywords, list)
                }
            if isinstance(data.get("subdomain_prefixes"), list):
                config["subdomain_prefixes"] = [
                    normalize_domain(str(prefix)).split(".")[0]
                    for prefix in data["subdomain_prefixes"]
                    if str(prefix).strip()
                ]
            profile_name = str(data.get("protection_profile", "default")).strip().lower()
            if profile_name in {"default", "child_protection"}:
                config["protection_profile"] = profile_name
            password_value = data.get("profile_password", "")
            stored_password = str(password_value) if password_value is not None else ""
            migrate_password = bool(stored_password) and not stored_password.startswith(PROFILE_PASSWORD_HASH_PREFIX)
            if stored_password.startswith(PROFILE_PASSWORD_HASH_PREFIX) and _parse_profile_password_hash(stored_password) is None:
                raise RuntimeError("The saved Safe Mode password hash is invalid. Restore a valid rules.json backup before continuing.")
            config["profile_password"] = hash_profile_password(stored_password) if migrate_password else stored_password
            if migrate_password:
                try:
                    save_rules(
                        config["blocked_domains"],
                        config["categories"],
                        config["subdomain_prefixes"],
                        config["protection_profile"],
                        config["profile_password"],
                    )
                except OSError as error:
                    raise RuntimeError("Avantis could not securely upgrade the saved Safe Mode password. Check that rules.json is writable, then restart the app.") from error
            return config
    except (OSError, json.JSONDecodeError, TypeError, ValueError):
        pass

    return default_rules()


def normalize_domain(value: str) -> str:
    value = value.strip().lower()
    if not value:
        return ""

    parsed = urlparse(value if "://" in value else f"//{value}")
    domain = parsed.netloc or parsed.path.split("/")[0]
    domain = domain.rsplit("@", 1)[-1].split(":", 1)[0]
    if domain.startswith("www."):
        domain = domain[4:]
    return domain.strip(".")


def is_valid_domain(domain: str) -> bool:
    if not domain or len(domain) > 253:
        return False
    try:
        ipaddress.ip_address(domain)
        return False
    except ValueError:
        return bool(VALID_DOMAIN_PATTERN.fullmatch(domain))


def is_protected_domain(domain: str) -> bool:
    return any(domain == protected or domain.endswith(f".{protected}") for protected in PROTECTED_DOMAINS)


def load_dns_blocklist(directory=None):
    root = Path(directory) if directory is not None else RESOURCE_ROOT
    domains = set()

    for filename in DNS_BLOCKLIST_FILENAMES:
        path = root / filename
        try:
            lines = path.read_text(encoding="utf-8", errors="ignore").splitlines()
        except OSError:
            continue

        for line in lines:
            content = line.split("#", 1)[0].strip()
            if not content:
                continue

            entries = re.split(r"[\s,;]+", content)
            try:
                ipaddress.ip_address(entries[0])
                entries = entries[1:]
            except ValueError:
                pass

            for entry in entries:
                labels = entry.split(".")
                if len(labels) > 4:
                    try:
                        address = ipaddress.ip_address(".".join(labels[:4]))
                    except ValueError:
                        address = None
                    if isinstance(address, ipaddress.IPv4Address):
                        entry = ".".join(labels[4:])

                domain = normalize_domain(entry)
                if is_valid_domain(domain) and not is_protected_domain(domain):
                    domains.add(domain)

    return domains


def keyword_matches(domain: str, keyword: str) -> bool:
    searchable = re.sub(r"[^a-z0-9]+", " ", normalize_domain(domain)).strip()
    search_keyword = re.sub(r"[^a-z0-9]+", " ", keyword.lower()).strip()
    if not searchable or not search_keyword:
        return False
    pattern = rf"(?<![a-z0-9]){re.escape(search_keyword)}(?![a-z0-9])"
    return re.search(pattern, searchable) is not None


def detect_categories(domain: str, categories=None):
    text = normalize_domain(domain)
    found = []
    for category, keywords in (categories or DEFAULT_CATEGORIES).items():
        if any(keyword_matches(text, keyword) for keyword in keywords):
            found.append(category)
    return found


def smart_signals(site: str, domain: str, blocklist):
    signals = []
    raw_value = site.strip()
    parsed_value = raw_value if "://" in raw_value else f"//{raw_value}"
    try:
        parsed = urlparse(parsed_value)
        hostname = parsed.hostname or domain
    except ValueError:
        parsed = None
        hostname = domain

    if "xn--" in domain or any(ord(character) > 127 for character in domain):
        signals.append((20, "The domain uses punycode or non-ASCII characters."))

    try:
        ipaddress.ip_address(hostname)
        signals.append((30, "The address uses a raw IP instead of a named domain."))
    except ValueError:
        pass

    if parsed and parsed.username:
        signals.append((25, "The URL contains a username before the host, which can hide the real destination."))

    if parsed and parsed.password:
        signals.append((25, "The URL contains embedded credentials."))

    if parsed and parsed.scheme and parsed.scheme.lower() == "http":
        signals.append((10, "The address does not use encrypted HTTPS."))

    if parsed and parsed.port not in (None, 80, 443):
        signals.append((10, "The URL uses an unusual network port."))

    if domain in SHORTENER_DOMAINS:
        signals.append((15, "This is a URL-shortening service, so the final destination is hidden."))

    labels = domain.split(".")
    if len(labels) >= 5:
        signals.append((10, "The hostname has many subdomain levels."))
    if len(domain) > 45:
        signals.append((10, "The hostname is unusually long."))
    if re.search(r"(?:\d[.-]?){4,}", domain):
        signals.append((10, "The hostname contains an unusual concentration of numbers."))

    if labels and labels[-1] in SUSPICIOUS_TLDS:
        signals.append((10, f"The top-level domain .{labels[-1]} is frequently used in disposable links."))

    for blocked in blocklist:
        blocked_labels = normalize_domain(blocked).split(".")
        if len(blocked_labels) < 2 or len(labels) < 2:
            continue
        candidate = ".".join(labels[-2:])
        target = ".".join(blocked_labels[-2:])
        similarity = difflib.SequenceMatcher(None, candidate, target).ratio()
        if candidate != target and similarity >= 0.86:
            signals.append((25, f"The domain closely resembles the protected domain {target}."))
            break

    return signals


def analyze_site(site: str, blocklist, categories, dns_blocklist=()):
    domain = normalize_domain(site)
    matched_keywords = []
    matched_categories = []

    for category, keywords in categories.items():
        category_matches = [keyword for keyword in keywords if keyword_matches(domain, keyword)]
        if category_matches:
            matched_categories.append(category)
            matched_keywords.extend(category_matches)

    is_blocked = any(domain == blocked or domain.endswith(f".{blocked}") for blocked in blocklist)
    dns_feed_blocked = any(domain == blocked or domain.endswith(f".{blocked}") for blocked in dns_blocklist)
    is_blocked = is_blocked or dns_feed_blocked
    signals = smart_signals(site, domain, blocklist) if domain else []
    score = min(100, len(set(matched_keywords)) * 15 + len(matched_categories) * 10 + (70 if is_blocked else 0) + sum(weight for weight, _ in signals))
    if score >= 70:
        level = "High risk"
    elif score >= 35:
        level = "Review"
    else:
        level = "Low risk"

    if is_blocked or level == "High risk":
        recommendation = "Block or avoid this site."
    elif level == "Review":
        recommendation = "Review the destination carefully before continuing."
    else:
        recommendation = "No strong risk signal was found by the local rules."

    return {
        "domain": domain,
        "blocked": is_blocked,
        "dns_feed_blocked": dns_feed_blocked,
        "categories": matched_categories,
        "keywords": list(dict.fromkeys(matched_keywords)),
        "signals": [message for _, message in signals],
        "score": score,
        "level": level,
        "recommendation": recommendation,
    }


def save_rules(blocklist, categories, subdomain_prefixes=None, protection_profile="default", profile_password=""):
    unique = []
    seen = set()
    for item in blocklist:
        domain = normalize_domain(item)
        if domain and not is_protected_domain(domain) and domain not in seen:
            unique.append(domain)
            seen.add(domain)

    profile_value = str(protection_profile or "default").strip().lower()
    if profile_value not in {"default", "child_protection"}:
        profile_value = "default"

    config = {
        "blocked_domains": unique,
        "categories": categories,
        "subdomain_prefixes": subdomain_prefixes or DEFAULT_SUBDOMAIN_PREFIXES,
        "protection_profile": profile_value,
        "profile_password": hash_profile_password(profile_password),
    }
    temporary_fd, temporary_path = tempfile.mkstemp(
        prefix=f"{RULES_FILE.name}.",
        suffix=".tmp",
        dir=str(RULES_FILE.parent),
    )
    temporary_file = Path(temporary_path)
    try:
        with os.fdopen(temporary_fd, "w", encoding="utf-8") as f:
            json.dump(config, f, indent=2)
            f.flush()
            os.fsync(f.fileno())
        os.replace(temporary_file, RULES_FILE)
    finally:
        if temporary_file.exists():
            temporary_file.unlink()

def load_activity():
    default = {"retention_hours": RETENTION_OPTIONS["24 hours"], "events": []}
    if not ACTIVITY_FILE.exists():
        return default

    try:
        with ACTIVITY_FILE.open("r", encoding="utf-8") as f:
            data = json.load(f)
        if not isinstance(data, dict):
            return default
        retention_hours = int(data.get("retention_hours", default["retention_hours"]))
        if retention_hours not in RETENTION_OPTIONS.values():
            retention_hours = default["retention_hours"]
        events = data.get("events", [])
        if not isinstance(events, list):
            events = []
    except (OSError, ValueError, TypeError, json.JSONDecodeError):
        return default

    cutoff = datetime.now(timezone.utc) - timedelta(hours=retention_hours)
    active_events = []
    for event in events:
        try:
            timestamp = datetime.fromisoformat(event["timestamp"])
            if timestamp >= cutoff:
                active_events.append(event)
        except (KeyError, TypeError, ValueError):
            continue

    result = {"retention_hours": retention_hours, "events": active_events}
    if len(active_events) != len(events):
        save_activity(active_events, retention_hours)
    return result


def save_activity(events, retention_hours):
    with ACTIVITY_FILE.open("w", encoding="utf-8") as f:
        json.dump({"retention_hours": retention_hours, "events": events[-500:]}, f, indent=2)


def record_activity(event_type, domain, detail):
    activity = load_activity()
    activity["events"].append({
        "timestamp": datetime.now(timezone.utc).isoformat(),
        "type": event_type,
        "domain": domain or "(unrecognized input)",
        "detail": detail,
    })
    save_activity(activity["events"], activity["retention_hours"])


def policy_snapshot(config):
    return {
        "blocked_domains": list(config.get("blocked_domains", [])),
        "categories": config.get("categories", {}),
        "subdomain_prefixes": list(config.get("subdomain_prefixes", DEFAULT_SUBDOMAIN_PREFIXES)),
    }


def load_policy_history():
    try:
        with POLICY_HISTORY_FILE.open("r", encoding="utf-8") as f:
            data = json.load(f)
        if isinstance(data, list):
            return [entry for entry in data if isinstance(entry, dict)][-50:]
    except (OSError, json.JSONDecodeError, TypeError):
        pass
    return []


def save_policy_history(entries):
    try:
        with POLICY_HISTORY_FILE.open("w", encoding="utf-8") as f:
            json.dump(entries[-50:], f, indent=2)
    except OSError:
        return False
    return True


def archive_policy(config, action):
    if not config:
        return False
    history = load_policy_history()
    timestamp = datetime.now(timezone.utc).isoformat()
    history.append({
        "id": timestamp,
        "timestamp": timestamp,
        "action": action,
        "policy": policy_snapshot(config),
    })
    return save_policy_history(history)


def load_admin_audit():
    try:
        with ADMIN_AUDIT_FILE.open("r", encoding="utf-8") as f:
            data = json.load(f)
        if isinstance(data, list):
            return [event for event in data if isinstance(event, dict)][-500:]
    except (OSError, json.JSONDecodeError, TypeError):
        pass
    return []


def record_admin_audit(action, detail):
    events = load_admin_audit()
    events.append({
        "timestamp": datetime.now(timezone.utc).isoformat(),
        "action": action,
        "detail": detail,
    })
    try:
        with ADMIN_AUDIT_FILE.open("w", encoding="utf-8") as f:
            json.dump(events[-500:], f, indent=2)
    except OSError:
        return False
    return True


def render_admin_report(report):
    def text(value):
        return html.escape(str(value if value is not None else ""))

    def timestamp(value):
        try:
            parsed = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
            if parsed.tzinfo is None:
                parsed = parsed.replace(tzinfo=timezone.utc)
            return parsed.astimezone().strftime("%Y-%m-%d %H:%M:%S %Z")
        except (TypeError, ValueError):
            return str(value or "Not recorded")

    def table(headers, rows, empty_message):
        heading = "".join(f"<th scope=\"col\">{text(header)}</th>" for header in headers)
        if not rows:
            content = f"<tr><td colspan=\"{len(headers)}\" class=\"empty\">{text(empty_message)}</td></tr>"
        else:
            content = "".join("<tr>" + "".join(f"<td>{text(value)}</td>" for value in row) + "</tr>" for row in rows)
        return f"<div class=\"table-wrap\"><table><thead><tr>{heading}</tr></thead><tbody>{content}</tbody></table></div>"

    policy = report.get("policy", {})
    categories = policy.get("categories", {})
    domains = policy.get("blocked_domains", [])
    prefixes = policy.get("subdomain_prefixes", [])
    protection_message, protection_level = report.get("hosts_status", ["Status unavailable", "warning"])
    activity = report.get("manual_check_activity", {})
    activity_events = activity.get("events", [])
    history = report.get("policy_history", [])
    audit = report.get("admin_audit", [])
    rule_count = sum(len(words) for words in categories.values() if isinstance(words, list))
    protection_message, status_label = protection_status_for_display(protection_message, protection_level)

    category_rows = [(category.replace("_", " ").title(), ", ".join(words)) for category, words in sorted(categories.items())]
    activity_rows = [(timestamp(event.get("timestamp")), event.get("domain"), event.get("detail")) for event in reversed(activity_events)]
    audit_rows = [(timestamp(event.get("timestamp")), format_audit_action(event.get("action", "")), format_audit_detail(event.get("detail", ""))) for event in reversed(audit)]
    history_rows = [(timestamp(version.get("timestamp")), version.get("action"), f"{version.get('domain_count', 0)} domains; {version.get('rule_count', 0)} detection rules") for version in reversed(history)]
    generated = timestamp(report.get("generated_at"))

    return f"""<!doctype html>
<html lang="en">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>Avantis FireWall Admin Report</title>
  <style>
    :root {{ color-scheme: light; font-family: 'Segoe UI', Arial, sans-serif; color: #244253; background: #f3f7f8; }}
    * {{ box-sizing: border-box; }}
    body {{ margin: 0; padding: 32px 20px; }}
    main {{ width: min(100%, 1040px); margin: 0 auto; }}
    header {{ padding: 25px 28px; border-top: 4px solid #119cad; background: #fff; }}
    .eyebrow {{ margin: 0 0 7px; color: #16899a; font-size: 11px; font-weight: 700; letter-spacing: 1.2px; }}
    h1 {{ margin: 0; color: #193b4d; font: 700 28px Georgia, serif; }}
    header p:last-child {{ margin: 8px 0 0; color: #6a8290; font-size: 13px; }}
    section {{ margin-top: 18px; padding: 20px 22px; border: 1px solid #dce8ec; background: #fff; }}
    h2 {{ margin: 0 0 14px; color: #23485c; font-size: 16px; }}
    .summary {{ display: grid; grid-template-columns: repeat(4, minmax(0, 1fr)); gap: 10px; }}
    .metric {{ min-width: 0; padding: 13px; border: 1px solid #e3ecef; background: #f8fbfc; }}
    .metric span, .metric strong {{ display: block; }}
    .metric span {{ color: #6b8390; font-size: 11px; }}
    .metric strong {{ margin-top: 6px; overflow-wrap: anywhere; color: #173f52; font-size: 17px; }}
    .note {{ margin: 12px 0 0; color: #536f7e; font-size: 12px; line-height: 1.6; }}
    .table-wrap {{ overflow-x: auto; }}
    table {{ width: 100%; border-collapse: collapse; text-align: left; }}
    th {{ padding: 9px 10px; color: #66808e; background: #f5f9fa; font-size: 10px; font-weight: 700; }}
    td {{ padding: 10px; border-top: 1px solid #e9eff1; color: #365665; font-size: 12px; line-height: 1.5; overflow-wrap: anywhere; }}
    .empty {{ color: #7b8d96; text-align: center; }}
    .prefixes {{ margin: 0; color: #536f7e; font-size: 12px; line-height: 1.7; overflow-wrap: anywhere; }}
    footer {{ padding: 17px 4px; color: #718792; font-size: 11px; }}
    @media (max-width: 700px) {{ body {{ padding: 14px 10px; }} section {{ padding: 16px 13px; }} .summary {{ grid-template-columns: repeat(2, minmax(0, 1fr)); }} h1 {{ font-size: 24px; }} }}
    @media print {{ body {{ padding: 0; background: #fff; }} section, header {{ break-inside: avoid; border-color: #ccd8dc; }} main {{ width: 100%; }} }}
  </style>
</head>
<body>
  <main>
    <header><p class="eyebrow">AVANTIS FIREWALL · ADMINISTRATION</p><h1>Protection report</h1><p>Generated {text(generated)}</p></header>
    <section><h2>At a glance</h2><div class="summary"><div class="metric"><span>Protection profile</span><strong>{text(policy.get('protection_profile', 'Unknown').replace('_', ' ').title())}</strong></div><div class="metric"><span>Blocked domains</span><strong>{len(domains)}</strong></div><div class="metric"><span>Detection rules</span><strong>{rule_count}</strong></div><div class="metric"><span>Device protection</span><strong>{text(status_label)}</strong></div></div><p class="note">{text(protection_message)}</p></section>
    <section><h2>Blocked domains</h2>{table(['Domain'], [(domain,) for domain in domains], 'No blocked domains are configured.')}</section>
    <section><h2>Detection rules</h2>{table(['Category', 'Configured terms'], category_rows, 'No detection categories are configured.')}</section>
    <section><h2>Subdomain coverage</h2><p class="prefixes">{text(', '.join(prefixes) if prefixes else 'No subdomain prefixes configured.')}</p></section>
    <section><h2>Manual website checks</h2><p class="note">These are checks entered manually in Avantis, not a browsing history. Retention: {text(activity.get('retention_hours', 24))} hours.</p>{table(['Time', 'Checked address', 'Result'], activity_rows, 'No manual checks are currently stored.')}</section>
    <section><h2>Administrative audit</h2>{table(['Time', 'Action', 'Details'], audit_rows, 'No administrative actions are recorded.')}</section>
    <section><h2>Policy history</h2>{table(['Time', 'Change', 'Saved policy'], history_rows, 'No previous policy versions are saved yet.')}</section>
    <footer>Generated locally by Avantis FireWall. This report does not contain the Safe Mode password. Review manual-check entries before sharing.</footer>
  </main>
</body>
</html>
"""


def protection_status_for_display(message, level):
    message = str(message or "")
    if re.search(r"\d+ Avantis hostnames in \d+ mappings\.", message):
        return "Website restrictions are active across this PC.", "Active"
    if "No Avantis-managed rules found" in message:
        return "Website restrictions are not currently applied.", "Not applied"
    if level == "error":
        return "Protection status could not be checked. Confirm Administrator approval and try again.", "Check needed"
    return "Protection status needs review. Confirm Administrator approval and check again.", "Review needed"


def format_audit_action(action):
    if action == "Hosts protection applied":
        return "Website protection applied"
    if action == "Hosts protection removed":
        return "Website protection removed"
    return action


def format_audit_detail(detail):
    detail = str(detail or "")
    if "blocked domains applied" in detail:
        return "Website rules were applied to this PC."
    if "Avantis hosts rules were removed" in detail:
        return "Avantis website restrictions were removed from this PC."
    if "No Avantis hosts rules were present" in detail:
        return "No Avantis website restrictions were active."
    return detail


def inspect_hosts_rules():
    if not os.path.exists(WINDOWS_HOSTS_PATH):
        return "Hosts file check failed: Windows hosts file not found.", "error"

    marker = "# >>> suspicious-site-blocker >>>"
    end_marker = "# <<< suspicious-site-blocker <<<"
    try:
        with open(WINDOWS_HOSTS_PATH, "r", encoding="utf-8") as hosts_file:
            content = hosts_file.read()
    except (OSError, UnicodeError) as error:
        return f"Hosts file check failed: {error}", "error"

    start_index = content.find(marker)
    end_index = content.find(end_marker, start_index + len(marker)) if start_index >= 0 else -1
    if start_index < 0 and end_index < 0:
        return "Hosts file check: No Avantis-managed rules found.", "success"
    if start_index < 0 or end_index < 0:
        return "Hosts file check: Incomplete Avantis markers; review the hosts file.", "warning"

    section = content[start_index + len(marker):end_index]
    hostnames = set()
    mapping_count = 0
    for line in section.splitlines():
        fields = line.split()
        if len(fields) >= 2 and not fields[0].startswith("#"):
            mapping_count += 1
            hostnames.update(hostname.lower() for hostname in fields[1:])

    return f"Hosts file check: {len(hostnames)} Avantis hostnames in {mapping_count} mappings.", "warning"


def write_hosts_file(blocklist, subdomain_prefixes=None):
    if not os.path.exists(WINDOWS_HOSTS_PATH):
        raise FileNotFoundError("Windows hosts file not found.")

    try:
        with open(WINDOWS_HOSTS_PATH, "r", encoding="utf-8") as f:
            content = f.read()
    except PermissionError as error:
           raise PermissionError(f"Windows denied read access to the hosts file: {error}. This can indicate file permissions or security software, not only missing administrator rights.") from error

    marker = "# >>> suspicious-site-blocker >>>"
    end_marker = "# <<< suspicious-site-blocker <<<"

    if marker in content and end_marker in content:
        start_index = content.index(marker)
        end_index = content.index(end_marker) + len(end_marker)
        content = content[:start_index] + content[end_index:]

    lines = []
    prefixes = subdomain_prefixes or DEFAULT_SUBDOMAIN_PREFIXES
    protected_hosts = PROTECTED_HOSTS | PROTECTED_DOMAINS
    for domain in blocklist:
        clean = normalize_domain(domain)
        if clean and not any(clean == protected or clean.endswith(f".{protected}") for protected in protected_hosts):
            hostnames = [clean] + [f"{prefix}.{clean}" for prefix in prefixes]
            for hostname in hostnames:
                lines.append(f"127.0.0.1 {hostname}")
                lines.append(f"0.0.0.0 {hostname}")

    block_section = f"\n{marker}\n" + "\n".join(lines) + f"\n{end_marker}\n"
    new_content = content.rstrip() + "\n" + block_section.strip() + "\n"

    try:
        if not HOSTS_BACKUP_PATH.exists():
            shutil.copy2(WINDOWS_HOSTS_PATH, HOSTS_BACKUP_PATH)
    except PermissionError as error:
        raise PermissionError(f"Windows denied the hosts-file backup at {HOSTS_BACKUP_PATH}: {error}") from error

    replace_windows_hosts_file(WINDOWS_HOSTS_PATH, new_content)


def remove_hosts_rules():
    if not os.path.exists(WINDOWS_HOSTS_PATH):
        raise FileNotFoundError("Windows hosts file not found.")

    marker = "# >>> suspicious-site-blocker >>>"
    end_marker = "# <<< suspicious-site-blocker <<<"
    try:
        with open(WINDOWS_HOSTS_PATH, "r", encoding="utf-8") as f:
            content = f.read()
    except PermissionError as error:
           raise PermissionError(f"Windows denied read access to the hosts file: {error}. This can indicate file permissions or security software, not only missing administrator rights.") from error

    if marker not in content or end_marker not in content:
        return False

    start_index = content.index(marker)
    end_index = content.index(end_marker) + len(end_marker)
    new_content = content[:start_index].rstrip() + content[end_index:].lstrip()
    replace_windows_hosts_file(WINDOWS_HOSTS_PATH, new_content)
    return True


def flush_dns_cache():
    try:
        result = subprocess.run(["ipconfig", "/flushdns"], capture_output=True, text=True, check=False)
    except OSError as error:
        return False, str(error)

    if result.returncode != 0:
        detail = (result.stderr or result.stdout).strip()
        return False, detail or f"ipconfig exited with status {result.returncode}"
    return True, ""


def is_admin():
    if os.name == "nt":
        import ctypes
        try:
            return bool(ctypes.windll.shell32.IsUserAnAdmin())
        except (AttributeError, OSError):
            return False

    try:
        return os.getuid() == 0
    except AttributeError:
        return False


def run_as_admin(operation):
    if os.name != "nt":
        return False

    import ctypes
    from ctypes import wintypes

    executable_path = Path(sys.executable).resolve()
    result_fd, result_path = tempfile.mkstemp(prefix="avantis-hosts-", suffix=".json")
    os.close(result_fd)

    if getattr(sys, "frozen", False):
        arguments = ["--hosts-helper", operation, result_path]
        working_directory = str(APP_ROOT)
    else:
        script_path = str(Path(__file__).resolve())
        arguments = [script_path, "--hosts-helper", operation, result_path]
        working_directory = str(APP_ROOT)

    class ShellExecuteInfo(ctypes.Structure):
        _fields_ = [
            ("cbSize", wintypes.DWORD),
            ("fMask", wintypes.ULONG),
            ("hwnd", wintypes.HWND),
            ("lpVerb", wintypes.LPCWSTR),
            ("lpFile", wintypes.LPCWSTR),
            ("lpParameters", wintypes.LPCWSTR),
            ("lpDirectory", wintypes.LPCWSTR),
            ("nShow", ctypes.c_int),
            ("hInstApp", ctypes.c_void_p),
            ("lpIDList", ctypes.c_void_p),
            ("lpClass", wintypes.LPCWSTR),
            ("hkeyClass", ctypes.c_void_p),
            ("dwHotKey", wintypes.DWORD),
            ("hIcon", ctypes.c_void_p),
            ("hProcess", wintypes.HANDLE),
        ]

    info = ShellExecuteInfo()
    info.cbSize = ctypes.sizeof(info)
    info.fMask = 0x40
    info.lpVerb = "runas"
    info.lpFile = str(executable_path)
    info.lpParameters = subprocess.list2cmdline(arguments)
    info.lpDirectory = working_directory
    info.nShow = 0

    shell32 = ctypes.WinDLL("shell32", use_last_error=True)
    shell32.ShellExecuteExW.argtypes = [ctypes.POINTER(ShellExecuteInfo)]
    shell32.ShellExecuteExW.restype = wintypes.BOOL
    kernel32 = ctypes.WinDLL("kernel32", use_last_error=True)
    kernel32.WaitForSingleObject.argtypes = [wintypes.HANDLE, wintypes.DWORD]
    kernel32.WaitForSingleObject.restype = wintypes.DWORD
    kernel32.GetExitCodeProcess.argtypes = [wintypes.HANDLE, ctypes.POINTER(wintypes.DWORD)]
    kernel32.GetExitCodeProcess.restype = wintypes.BOOL
    kernel32.CloseHandle.argtypes = [wintypes.HANDLE]
    kernel32.CloseHandle.restype = wintypes.BOOL
    try:
        if not shell32.ShellExecuteExW(ctypes.byref(info)):
            error_code = ctypes.get_last_error()
            if error_code == 1223:
                raise PermissionError("Administrator approval was cancelled.")
            raise OSError(error_code, ctypes.FormatError(error_code))

        if not info.hProcess:
            raise OSError("Windows did not provide a process handle for the elevated operation.")
        wait_result = kernel32.WaitForSingleObject(info.hProcess, 0xFFFFFFFF)
        if wait_result != 0:
            raise ctypes.WinError(ctypes.get_last_error())

        exit_code = wintypes.DWORD()
        if not kernel32.GetExitCodeProcess(info.hProcess, ctypes.byref(exit_code)):
            raise ctypes.WinError(ctypes.get_last_error())

        response_path = Path(result_path)
        if not response_path.is_file():
            raise RuntimeError(f"The elevated hosts operation exited with status {exit_code.value} without returning a result.")
        response = json.loads(response_path.read_text(encoding="utf-8"))
        if "error" in response:
            raise RuntimeError(response["error"])
        if exit_code.value != 0 or "result" not in response:
            raise RuntimeError(f"The elevated hosts operation failed with status {exit_code.value}.")
        return response["result"]
    finally:
        if info.hProcess:
            kernel32.CloseHandle(info.hProcess)
        try:
            os.unlink(result_path)
        except FileNotFoundError:
            pass


def run_hosts_helper(operation, result_path):
    try:
        if operation not in {"apply", "remove"}:
            raise ValueError("Unknown elevated hosts operation.")
        if not is_admin():
            raise PermissionError("The elevated hosts helper did not receive administrator access.")
        initialize_data_root()
        api = DesktopApi()
        result = api.apply_hosts() if operation == "apply" else api.remove_hosts()
        response = {"result": result}
    except Exception as error:
        response = {"error": str(error)}
        exit_code = 1
    else:
        exit_code = 0

    Path(result_path).write_text(json.dumps(response), encoding="utf-8")
    return exit_code


def replace_windows_hosts_file(target_path: str, new_content: str):
    temp_path = f"{target_path}.avantis.tmp"
    try:
        with open(temp_path, "w", encoding="utf-8", newline="\n") as f:
            f.write(new_content)
    except PermissionError as error:
        raise PermissionError(f"Windows denied writing the temporary hosts file at {temp_path}: {error}") from error

    try:
        try:
            os.replace(temp_path, target_path)
            return
        except PermissionError:
            pass

        import ctypes
        if os.name == "nt":
            kernel32 = ctypes.windll.kernel32
            if kernel32.MoveFileExW(temp_path, target_path, 0x1):
                return

        raise PermissionError(f"Windows denied replacing the hosts file at {target_path}. Check file permissions or security software that may be locking the file.")
    finally:
        try:
            if os.path.exists(temp_path):
                os.remove(temp_path)
        except OSError:
            pass


def show_app_dialog(parent, title, message, kind="info", confirm=False):
    palette = {
        "info": ("#087f92", "Information"),
        "success": ("#16805c", "Success"),
        "warning": ("#b7791f", "Review"),
        "error": ("#c24141", "Action required"),
        "risk": ("#b42318", "RISK FLAGGED"),
    }
    accent, label = palette[kind]
    result = {"value": False}
    dialog = tk.Toplevel(parent)
    dialog.title(title)
    dialog.configure(bg="#eaf4f7")
    dialog.resizable(False, False)
    dialog.transient(parent)
    if hasattr(parent, "window_icon"):
        dialog.iconphoto(True, parent.window_icon)

    header_bg = accent if kind == "risk" else "#0b6478"
    header = tk.Frame(dialog, bg=header_bg, padx=20, pady=14)
    header.pack(fill="x")
    if kind == "risk":
        risk_row = tk.Frame(header, bg=header_bg)
        risk_row.pack(anchor="w")
        tk.Label(risk_row, text="⚠", bg=header_bg, fg="#ffffff", font=("Segoe UI Symbol", 18, "bold")).pack(side="left", padx=(0, 7))
        tk.Label(risk_row, text=label, bg=header_bg, fg="#ffffff", font=("Segoe UI", 10, "bold")).pack(side="left")
    else:
        tk.Label(header, text=label, bg=header_bg, fg="#ffffff", font=("Segoe UI", 10, "bold")).pack(anchor="w")
    tk.Label(header, text=title, bg=header_bg, fg="#ffffff", font=("Segoe UI", 16, "bold")).pack(anchor="w", pady=(2, 0))

    body = tk.Frame(dialog, bg="#ffffff", padx=20, pady=18)
    body.pack(fill="both", expand=True)
    tk.Label(body, text=message, justify="left", anchor="w", wraplength=460, bg="#ffffff", fg="#16445a", font=("Segoe UI", 10)).pack(fill="x")

    actions = tk.Frame(body, bg="#ffffff")
    actions.pack(fill="x", pady=(18, 0))

    def close(value=False):
        result["value"] = value
        dialog.destroy()

    if confirm:
        tk.Button(actions, text="Cancel", command=lambda: close(False), bg="#ffffff", fg="#16445a", activebackground="#dceff3", relief="solid", borderwidth=1, padx=14, pady=6).pack(side="right", padx=(8, 0))
        tk.Button(actions, text="Continue", command=lambda: close(True), bg=accent, fg="#ffffff", activebackground="#075d70", relief="flat", borderwidth=0, padx=14, pady=7).pack(side="right")
    else:
        tk.Button(actions, text="Close", command=close, bg=accent, fg="#ffffff", activebackground="#075d70", relief="flat", borderwidth=0, padx=18, pady=7).pack(side="right")

    dialog.protocol("WM_DELETE_WINDOW", close)
    dialog.update_idletasks()
    parent.update_idletasks()
    position_x = parent.winfo_rootx() + (parent.winfo_width() - dialog.winfo_width()) // 2
    position_y = parent.winfo_rooty() + (parent.winfo_height() - dialog.winfo_height()) // 2
    dialog.geometry(f"+{max(position_x, 0)}+{max(position_y, 0)}")
    dialog.grab_set()
    dialog.focus_set()
    parent.wait_window(dialog)
    return result["value"]


class BlockerApp(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Avantis FireWall")
        self.geometry("900x680")
        self.minsize(760, 580)
        self.configure(bg="#eaf4f7")

        style = ttk.Style(self)
        style.theme_use("clam")
        style.configure("App.TFrame", background="#eaf4f7")
        style.configure("App.TLabel", background="#eaf4f7", foreground="#16445a")
        style.configure("Header.TFrame", background="#0b6478")
        style.configure("Header.TLabel", background="#0b6478", foreground="#ffffff")
        style.configure("Panel.TFrame", background="#ffffff")
        style.configure("Panel.TLabel", background="#ffffff", foreground="#12445b")
        style.configure("Muted.TLabel", background="#ffffff", foreground="#527184")
        style.configure("Accent.TButton", background="#087f92", foreground="#ffffff", padding=(12, 7))
        style.map("Accent.TButton", background=[("active", "#075d70")])
        style.configure("Secondary.TButton", background="#dcebee", foreground="#16445a", padding=(12, 8), borderwidth=0, relief="flat")
        style.map("Secondary.TButton", background=[("active", "#c4e1e5"), ("pressed", "#b7d9de")], foreground=[("disabled", "#8aa0a8")])
        style.configure("Accent.TButton", background="#087f92", foreground="#ffffff", padding=(12, 8), borderwidth=0, relief="flat")
        style.map("Accent.TButton", background=[("active", "#075d70"), ("pressed", "#064b5a")], foreground=[("disabled", "#d0e5e8")])
        style.configure("TNotebook", background="#eaf4f7", borderwidth=0, tabmargins=(0, 0, 0, 0))
        style.configure("TNotebook.Tab", background="#d4e5e8", foreground="#16445a", padding=(20, 10), borderwidth=0, font=("Segoe UI", 10, "bold"))
        style.map("TNotebook.Tab", background=[("selected", "#ffffff"), ("active", "#c4e1e5")], foreground=[("selected", "#075d70")])
        style.configure("Card.TFrame", background="#ffffff")
        style.configure("CardValue.TLabel", background="#ffffff", foreground="#0b6478", font=("Segoe UI", 22, "bold"))
        style.configure("CardTitle.TLabel", background="#ffffff", foreground="#527184", font=("Segoe UI", 9, "bold"))
        style.configure("CardNote.TLabel", background="#ffffff", foreground="#7b93a0", font=("Segoe UI", 9))
        style.configure("SecurityCard.TFrame", background="#eaf4f7")
        style.configure("SecureBadge.TLabel", background="#0b6478", foreground="#ffffff", font=("Segoe UI", 11, "bold"))

        config = load_config()
        self.blocklist = [normalize_domain(item) for item in config["blocked_domains"] if str(item).strip()]
        self.categories = config["categories"]
        self.subdomain_prefixes = config["subdomain_prefixes"]
        self.current_profile = config.get("protection_profile", "default")
        self.profile_password = str(config.get("profile_password", "") or "")
        self.create_widgets()
        self.refresh_listbox()
        self.refresh_hosts_status()
        self.bind_all("<Control-Shift-P>", self.handle_admin_access)

    def create_widgets(self):
        header = ttk.Frame(self, style="Header.TFrame", padding=(26, 20))
        header.pack(fill="x")
        brand_row = ttk.Frame(header, style="Header.TFrame")
        brand_row.pack(fill="x")
        logo = tk.Canvas(brand_row, width=270, height=62, bg="#ffffff", highlightthickness=0)
        logo.pack(side="left", padx=(0, 20))
        if LOGO_FILE.exists():
            self.logo_image = tk.PhotoImage(file=str(LOGO_FILE))
            source_logo = Image.open(LOGO_FILE)
            icon_size = min(source_logo.width, source_logo.height)
            self.window_icon = ImageTk.PhotoImage(source_logo.crop((0, 0, icon_size, icon_size)))
            scale = max(1, (self.logo_image.width() + 269) // 270, (self.logo_image.height() + 61) // 62)
            if scale > 1:
                self.logo_image = self.logo_image.subsample(scale, scale)
            logo.create_image(135, 31, image=self.logo_image)
            self.iconphoto(True, self.window_icon)
        title_block = ttk.Frame(brand_row, style="Header.TFrame")
        title_block.pack(side="left", fill="x", expand=True)
        ttk.Label(title_block, text="Avantis FireWall", style="Header.TLabel", font=("Segoe UI", 21, "bold")).pack(anchor="w")
        ttk.Label(title_block, text="Manage safer browsing with simple, editable rules", style="Header.TLabel", font=("Segoe UI", 10)).pack(anchor="w", pady=(4, 0))

        content = ttk.Frame(self, style="App.TFrame", padding=(18, 14))
        content.pack(fill="both", expand=True)

        self.notebook = ttk.Notebook(content)
        self.notebook.pack(fill="both", expand=True)
        notebook = self.notebook

        protect_tab = ttk.Frame(notebook, style="App.TFrame", padding=14)
        import_tab = ttk.Frame(notebook, style="App.TFrame", padding=14)
        domains_tab = ttk.Frame(notebook, style="App.TFrame", padding=14)
        insights_tab = ttk.Frame(notebook, style="App.TFrame", padding=14)
        notebook.add(protect_tab, text="Protect")
        notebook.add(import_tab, text="Import")
        notebook.add(domains_tab, text="Domains")
        notebook.add(insights_tab, text="Insights")

        profile_panel = ttk.Frame(protect_tab, style="Panel.TFrame", padding=(18, 18, 18, 14))
        profile_panel.pack(fill="x", pady=(0, 12))

        mode_header = ttk.Frame(profile_panel, style="Panel.TFrame")
        mode_header.pack(fill="x")
        ttk.Label(mode_header, text="Protection mode", style="Panel.TLabel", font=("Segoe UI", 11, "bold")).pack(anchor="w")
        self.profile_status_label = ttk.Label(mode_header, text="", style="Muted.TLabel")
        self.profile_status_label.pack(anchor="e", pady=(0, 6))

        mode_card = tk.Frame(profile_panel, bg="#edf8fb", highlightbackground="#a9d8e4", highlightthickness=1, padx=18, pady=18)
        mode_card.pack(fill="x")

        mode_left = tk.Frame(mode_card, bg="#edf8fb")
        mode_left.pack(side="left", fill="x", expand=True)
        self.mode_lock_icon = tk.Label(mode_left, text="🔒", bg="#edf8fb", fg="#0b6478", font=("Segoe UI", 22, "bold"))
        self.mode_lock_icon.pack(anchor="w", pady=(0, 4))
        self.current_profile_label = ttk.Label(mode_left, text="Current profile", style="Panel.TLabel", font=("Segoe UI", 10, "bold"))
        self.current_profile_label.pack(anchor="w")
        self.profile_var = tk.StringVar(value=PROFILE_LABELS[self.current_profile])
        self.profile_combo = ttk.Combobox(mode_left, textvariable=self.profile_var, values=list(PROFILE_VALUES), state="readonly", width=22)
        self.profile_combo.pack(anchor="w", pady=(6, 0))
        self.profile_combo.bind("<<ComboboxSelected>>", self.handle_profile_change)

        mode_right = tk.Frame(mode_card, bg="#edf8fb")
        mode_right.pack(side="right", anchor="center")
        self.profile_badge = tk.Label(mode_right, text="Secure mode", bg="#0b6478", fg="#ffffff", font=("Segoe UI", 11, "bold"), padx=18, pady=8)
        self.profile_badge.pack(anchor="e")
        self.set_password_button = ttk.Button(mode_right, text="Set Safe Mode password", command=self.set_profile_password, style="Secondary.TButton")
        self.set_password_button.pack(anchor="e", pady=(8, 0))
        self.admin_access_button = ttk.Button(mode_right, text="Admin access", command=self.open_admin_access, style="Accent.TButton")
        self.admin_access_button.pack(anchor="e", pady=(6, 0))
        self.refresh_profile_status()

        self.input_frame = ttk.Frame(protect_tab, style="Panel.TFrame", padding=18)
        input_frame = self.input_frame
        input_frame.pack(fill="x", pady=(0, 12))
        input_frame.columnconfigure(0, weight=3)
        input_frame.columnconfigure(1, weight=2)

        self.analyze_panel = ttk.Frame(input_frame, style="Panel.TFrame", padding=(0, 0, 14, 0))
        analyze_panel = self.analyze_panel
        analyze_panel.grid(row=0, column=0, sticky="nsew")
        ttk.Label(analyze_panel, text="Analyze a site", style="Panel.TLabel", font=("Segoe UI", 11, "bold")).pack(anchor="w")
        ttk.Label(analyze_panel, text="Check the risk of a URL or domain", style="Muted.TLabel").pack(anchor="w", pady=(2, 6))
        self.check_var = tk.StringVar()
        check_entry = ttk.Entry(analyze_panel, textvariable=self.check_var, font=("Segoe UI", 10))
        check_entry.pack(fill="x", pady=(0, 8))
        ttk.Button(analyze_panel, text="Analyze Risk", command=self.check_url, style="Secondary.TButton").pack(anchor="w")

        self.add_panel = ttk.Frame(input_frame, style="Panel.TFrame", padding=(14, 0, 0, 0))
        add_panel = self.add_panel
        add_panel.grid(row=0, column=1, sticky="nsew")
        ttk.Label(add_panel, text="Add a blocked domain", style="Panel.TLabel", font=("Segoe UI", 11, "bold")).pack(anchor="w")
        ttk.Label(add_panel, text="Add a domain to the local blocklist", style="Muted.TLabel").pack(anchor="w", pady=(2, 6))
        self.add_var = tk.StringVar()
        add_entry = ttk.Entry(add_panel, textvariable=self.add_var, font=("Segoe UI", 10))
        add_entry.pack(fill="x", pady=(0, 8))
        ttk.Button(add_panel, text="Add Domain", command=self.add_url, style="Accent.TButton").pack(anchor="w")

        self.host_frame = ttk.Frame(protect_tab, style="Panel.TFrame", padding=18)
        host_frame = self.host_frame
        host_frame.pack(fill="x")
        ttk.Label(host_frame, text="Protection controls", style="Panel.TLabel", font=("Segoe UI", 11, "bold")).pack(anchor="w")
        ttk.Label(host_frame, text="Apply or remove Avantis website restrictions. Administrator approval is required.", style="Muted.TLabel").pack(anchor="w", pady=(2, 10))
        host_buttons = ttk.Frame(host_frame, style="Panel.TFrame")
        host_buttons.pack(fill="x")
        ttk.Button(host_buttons, text="Apply Protection", command=self.apply_blocking, style="Accent.TButton").pack(side="left", padx=(0, 8))
        ttk.Button(host_buttons, text="Remove Protection", command=self.remove_blocking, style="Secondary.TButton").pack(side="left", padx=(0, 8))
        ttk.Button(host_buttons, text="Check Status", command=self.refresh_hosts_status, style="Secondary.TButton").pack(side="left")
        self.hosts_status_label = ttk.Label(host_frame, text="Protection status has not been checked.", style="Muted.TLabel", wraplength=800)
        self.hosts_status_label.pack(anchor="w", pady=(10, 0))

        self.overview = ttk.Frame(protect_tab, style="App.TFrame")
        overview = self.overview
        overview.pack(fill="both", expand=True, pady=(12, 0))
        ttk.Label(overview, text="Protection overview", style="App.TLabel", font=("Segoe UI", 12, "bold")).pack(anchor="w", pady=(0, 8))
        ttk.Label(overview, text="A quiet view of what Avantis is ready to protect on this computer.", style="App.TLabel").pack(anchor="w", pady=(0, 10))
        cards = ttk.Frame(overview, style="App.TFrame")
        cards.pack(fill="x")
        for column in range(3):
            cards.columnconfigure(column, weight=1)

        domain_card = ttk.Frame(cards, style="Card.TFrame", padding=16)
        domain_card.grid(row=0, column=0, sticky="nsew", padx=(0, 8))
        self.domain_stat_label = ttk.Label(domain_card, text="0", style="CardValue.TLabel")
        self.domain_stat_label.pack(anchor="w")
        ttk.Label(domain_card, text="BLOCKED DOMAINS", style="CardTitle.TLabel").pack(anchor="w", pady=(4, 0))
        ttk.Label(domain_card, text="Ready to apply website protection", style="CardNote.TLabel").pack(anchor="w", pady=(2, 0))

        rule_card = ttk.Frame(cards, style="Card.TFrame", padding=16)
        rule_card.grid(row=0, column=1, sticky="nsew", padx=8)
        self.rule_stat_label = ttk.Label(rule_card, text="0", style="CardValue.TLabel")
        self.rule_stat_label.pack(anchor="w")
        ttk.Label(rule_card, text="DETECTION WORDS", style="CardTitle.TLabel").pack(anchor="w", pady=(4, 0))
        ttk.Label(rule_card, text="Local risk dictionary", style="CardNote.TLabel").pack(anchor="w", pady=(2, 0))

        mode_card = ttk.Frame(cards, style="Card.TFrame", padding=16)
        mode_card.grid(row=0, column=2, sticky="nsew", padx=(8, 0))
        ttk.Label(mode_card, text="LOCAL", style="CardValue.TLabel").pack(anchor="w")
        ttk.Label(mode_card, text="PRIVACY MODE", style="CardTitle.TLabel").pack(anchor="w", pady=(4, 0))
        ttk.Label(mode_card, text="No browsing data leaves this app", style="CardNote.TLabel").pack(anchor="w", pady=(2, 0))

        bulk_frame = ttk.Frame(import_tab, style="Panel.TFrame", padding=18)
        bulk_frame.pack(fill="both", expand=True)
        ttk.Label(bulk_frame, text="Bulk import", style="Panel.TLabel", font=("Segoe UI", 11, "bold")).pack(anchor="w")
        ttk.Label(bulk_frame, text="Paste multiple domains separated by new lines, spaces, or commas", style="Muted.TLabel").pack(anchor="w", pady=(2, 6))

        bulk_row = ttk.Frame(bulk_frame, style="Panel.TFrame")
        bulk_row.pack(fill="x")
        self.bulk_text = tk.Text(bulk_row, height=5, wrap="word", font=("Segoe UI", 10), relief="flat", borderwidth=1, highlightthickness=1, highlightbackground="#cbd5e1", highlightcolor="#087f92")
        self.bulk_text.pack(side="left", fill="both", expand=True)
        bulk_actions = ttk.Frame(bulk_row, style="Panel.TFrame")
        bulk_actions.pack(side="left", padx=(12, 0), anchor="n")
        ttk.Button(bulk_actions, text="Add All", command=self.add_multiple, style="Accent.TButton").pack(fill="x")
        ttk.Button(bulk_actions, text="Import File", command=self.import_file, style="Secondary.TButton").pack(fill="x", pady=(8, 0))

        def resize_bulk_input(_event=None):
            if not self.bulk_text.edit_modified():
                return
            line_count = int(self.bulk_text.index("end-1c").split(".")[0])
            self.bulk_text.configure(height=min(12, max(5, line_count)))
            self.bulk_text.edit_modified(False)

        self.bulk_text.bind("<<Modified>>", resize_bulk_input)
        ttk.Label(bulk_frame, text="Imported domains are saved locally. Website protection changes only when you select Apply Protection.", style="Muted.TLabel").pack(anchor="w", pady=(10, 0))

        list_frame = ttk.Frame(domains_tab, style="Panel.TFrame", padding=18)
        list_frame.pack(fill="both", expand=True)
        list_header = ttk.Frame(list_frame, style="Panel.TFrame")
        list_header.pack(fill="x", pady=(0, 8))
        ttk.Label(list_header, text="Blocked domains", style="Panel.TLabel", font=("Segoe UI", 11, "bold")).pack(side="left")
        self.count_label = ttk.Label(list_header, text="", style="Muted.TLabel")
        self.count_label.pack(side="right")

        list_area = ttk.Frame(list_frame, style="Panel.TFrame")
        list_area.pack(fill="both", expand=True)
        scrollbar = ttk.Scrollbar(list_area, orient="vertical")
        scrollbar.pack(side="right", fill="y")
        self.listbox = tk.Listbox(list_area, height=12, activestyle="none", selectmode=tk.EXTENDED, font=("Segoe UI", 10), bg="#f8fafc", fg="#16324f", selectbackground="#087f92", selectforeground="#ffffff", relief="flat", borderwidth=0, yscrollcommand=scrollbar.set)
        self.listbox.pack(side="left", fill="both", expand=True)
        scrollbar.config(command=self.listbox.yview)

        controls = ttk.Frame(list_frame, style="Panel.TFrame")
        controls.pack(fill="x", pady=(10, 0))
        ttk.Button(controls, text="Remove Selected", command=self.remove_selected, style="Secondary.TButton").pack(side="left", padx=(0, 8))
        ttk.Button(controls, text="Clear All", command=self.clear_all, style="Secondary.TButton").pack(side="left")

        insights_panel = ttk.Frame(insights_tab, style="App.TFrame", padding=(4, 8))
        insights_panel.pack(fill="both", expand=True)
        insights_header = ttk.Frame(insights_panel, style="App.TFrame")
        insights_header.pack(fill="x")
        ttk.Label(insights_header, text="Insights & tools", style="App.TLabel", font=("Segoe UI", 16, "bold")).pack(anchor="w")
        ttk.Label(insights_header, text="Review activity, look up protection data, and manage detection rules.", style="App.TLabel").pack(anchor="w", pady=(4, 0))
        tk.Frame(insights_panel, bg="#b7dfe6", height=2).pack(fill="x", pady=(16, 18))

        insight_cards = ttk.Frame(insights_panel, style="App.TFrame")
        insight_cards.pack(fill="x")
        for column in range(3):
            insight_cards.columnconfigure(column, weight=1, uniform="insight-card")

        insight_tools = (
            ("01", "ACTIVITY", "Safety Center", "Review locally stored records of manual site checks.", "Review activity", self.open_safety_center, "#087f92"),
            ("02", "LOOKUP", "Site Dictionary", "Search blocked domains, detection words, and hostname prefixes.", "Search entries", self.open_site_dictionary, "#b7791f"),
            ("03", "CONFIGURE", "Rule Library", "Add or remove words used by site risk analysis.", "Manage rules", self.open_rule_manager, "#2b7a78"),
        )
        for column, (number, category, title, description, action, command, accent) in enumerate(insight_tools):
            outer_pad = (0, 8) if column == 0 else ((8, 8) if column == 1 else (8, 0))
            card = tk.Frame(insight_cards, bg="#ffffff", highlightbackground="#c9dce2", highlightthickness=1, padx=0, pady=0)
            card.grid(row=0, column=column, sticky="nsew", padx=outer_pad)
            tk.Frame(card, bg=accent, height=4).pack(fill="x")
            card_body = tk.Frame(card, bg="#ffffff", padx=18, pady=17)
            card_body.pack(fill="both", expand=True)
            top_line = tk.Frame(card_body, bg="#ffffff")
            top_line.pack(fill="x")
            tk.Label(top_line, text=number, bg=accent, fg="#ffffff", font=("Segoe UI", 10, "bold"), padx=9, pady=5).pack(side="left")
            tk.Label(top_line, text=category, bg="#ffffff", fg=accent, font=("Segoe UI", 9, "bold")).pack(side="left", padx=(10, 0))
            tk.Label(card_body, text=title, bg="#ffffff", fg="#163b4d", font=("Segoe UI", 14, "bold")).pack(anchor="w", pady=(17, 5))
            tk.Label(card_body, text=description, bg="#ffffff", fg="#5f7f8f", font=("Segoe UI", 9), justify="left", anchor="w", wraplength=235, height=3).pack(fill="x")
            ttk.Button(card_body, text=action, command=command, style="Secondary.TButton").pack(anchor="w", pady=(18, 0))

        self.restricted_tabs = (import_tab, domains_tab, insights_tab)
        self.refresh_profile_visibility()

    def save_current_rules(self):
        save_rules(self.blocklist, self.categories, self.subdomain_prefixes, self.current_profile, self.profile_password)

    def refresh_profile_status(self):
        if self.current_profile == "child_protection":
            if self.profile_password:
                self.profile_status_label.config(text="Locked · admin password required")
                badge_text = "SAFE MODE"
                badge_bg = "#0b6478"
            else:
                self.profile_status_label.config(text="Password not set")
                badge_text = "SAFE MODE"
                badge_bg = "#b7791f"
        else:
            self.profile_status_label.config(text="Editing enabled")
            badge_text = "ADMIN"
            badge_bg = "#2b7a78"

        if hasattr(self, "profile_badge"):
            self.profile_badge.config(text=badge_text, bg=badge_bg)
        if hasattr(self, "mode_lock_icon"):
            self.mode_lock_icon.config(text="🔒" if self.current_profile == "child_protection" else "🛡️")

        self.refresh_profile_visibility()

    def refresh_profile_visibility(self):
        if not hasattr(self, "restricted_tabs"):
            return

        child_mode = self.current_profile == "child_protection"
        tab_state = "hidden" if child_mode else "normal"
        for tab in self.restricted_tabs:
            self.notebook.tab(tab, state=tab_state)

        if child_mode:
            self.admin_access_button.config(text="Admin access")
            self.current_profile_label.pack_forget()
            self.profile_combo.pack_forget()
            self.set_password_button.pack_forget()
            self.admin_access_button.pack(anchor="e", pady=(8, 0))
            self.add_panel.grid_remove()
            self.analyze_panel.grid_configure(columnspan=2, padx=0)
            self.input_frame.columnconfigure(0, weight=1)
            self.input_frame.columnconfigure(1, weight=0)
            self.host_frame.pack_forget()
            self.overview.pack_forget()
        else:
            self.admin_access_button.pack_forget()
            self.current_profile_label.pack(anchor="w")
            self.profile_combo.pack(anchor="w", pady=(6, 0))
            self.set_password_button.pack(anchor="e", pady=(8, 0))
            self.analyze_panel.grid_configure(columnspan=1, padx=(0, 14))
            self.input_frame.columnconfigure(0, weight=3)
            self.input_frame.columnconfigure(1, weight=2)
            self.add_panel.grid()
            self.host_frame.pack(fill="x")
            self.overview.pack(fill="both", expand=True, pady=(12, 0))

    def handle_admin_access(self, _event=None):
        if self.current_profile == "child_protection":
            self.open_admin_access()
        return "break"

    def open_admin_access(self):
        if self.current_profile == "child_protection":
            password = simpledialog.askstring(
                "Admin access",
                "Enter the admin password to open Safe Mode controls:",
                show="*",
                parent=self,
            )
            if not verify_profile_password(password, self.profile_password):
                show_app_dialog(self, "Access blocked", "The admin password is incorrect.", "warning")
                return

        dashboard = tk.Toplevel(self)
        dashboard.title("Admin access")
        dashboard.configure(bg="#eaf4f7")
        dashboard.geometry("520x420")
        dashboard.resizable(False, False)
        dashboard.transient(self)
        dashboard.grab_set()

        if hasattr(self, "window_icon"):
            dashboard.iconphoto(True, self.window_icon)

        header = tk.Frame(dashboard, bg="#0b6478", padx=20, pady=18)
        header.pack(fill="x")
        tk.Label(header, text="🔒 Admin access", bg="#0b6478", fg="#ffffff", font=("Segoe UI", 18, "bold")).pack(anchor="w")
        tk.Label(header, text="Review Safe Mode and manage access controls", bg="#0b6478", fg="#d9f3f7", font=("Segoe UI", 10)).pack(anchor="w", pady=(4, 0))

        body = tk.Frame(dashboard, bg="#ffffff", padx=20, pady=20)
        body.pack(fill="both", expand=True)

        summary = tk.Frame(body, bg="#ffffff")
        summary.pack(fill="x")
        stats = [
            ("Profile", PROFILE_LABELS[self.current_profile]),
            ("Password", "Enabled" if self.profile_password else "Disabled"),
            ("Blocked", str(len(self.blocklist))),
        ]
        for idx, (label, value) in enumerate(stats):
            card = tk.Frame(summary, bg="#edf8fb", highlightbackground="#d0e1e7", highlightthickness=1, padx=12, pady=12)
            card.grid(row=0, column=idx, sticky="nsew", padx=(0 if idx == 0 else 8, 0))
            tk.Label(card, text=label, bg="#edf8fb", fg="#58727d", font=("Segoe UI", 9, "bold")).pack(anchor="w")
            tk.Label(card, text=value, bg="#edf8fb", fg="#0f3140", font=("Segoe UI", 18, "bold")).pack(anchor="w", pady=(4, 0))
        for column_index in (0, 1, 2):
            summary.columnconfigure(column_index, weight=1)

        actions = tk.Frame(body, bg="#ffffff")
        actions.pack(fill="x", pady=(20, 0))
        tk.Label(actions, text="Controls", bg="#ffffff", fg="#12445b", font=("Segoe UI", 11, "bold")).pack(anchor="w")

        btn_frame = tk.Frame(actions, bg="#ffffff")
        btn_frame.pack(fill="x", pady=(8, 0))
        ttk.Button(btn_frame, text="Set / change password", command=lambda: (dashboard.destroy(), self.set_profile_password()), style="Secondary.TButton").pack(side="left", padx=(0, 8))
        if self.current_profile == "child_protection":
            ttk.Button(btn_frame, text="Restore Admin", command=lambda: self.restore_admin_profile(dashboard), style="Accent.TButton").pack(side="left")
        else:
            ttk.Button(btn_frame, text="Switch profile", command=lambda: (dashboard.destroy(), self.profile_combo.focus_set()), style="Accent.TButton").pack(side="left")

        note = tk.Label(body, text="Only the admin password can turn off Safe Mode and restore Admin mode.", bg="#ffffff", fg="#496573", justify="left", wraplength=440, font=("Segoe UI", 9))
        note.pack(anchor="w", pady=(18, 0))

        footer = tk.Frame(body, bg="#ffffff")
        footer.pack(fill="x", pady=(18, 0))
        ttk.Button(footer, text="Close", command=dashboard.destroy, style="Secondary.TButton").pack(anchor="e")

        dashboard.update_idletasks()
        self.update_idletasks()
        position_x = self.winfo_x() + (self.winfo_width() - dashboard.winfo_width()) // 2
        position_y = self.winfo_y() + (self.winfo_height() - dashboard.winfo_height()) // 2
        dashboard.geometry(f"+{max(position_x, 0)}+{max(position_y, 0)}")

    def restore_admin_profile(self, dashboard):
        self.current_profile = "default"
        self.profile_var.set(PROFILE_LABELS[self.current_profile])
        self.save_current_rules()
        self.refresh_profile_status()
        dashboard.destroy()

    def set_profile_password(self):
        if self.profile_password:
            current = simpledialog.askstring(
                "Current password",
                "Enter the current admin password to change it:",
                show="*",
                parent=self,
            )
            if not verify_profile_password(current, self.profile_password):
                show_app_dialog(self, "Password change denied", "The current password is incorrect.", "warning")
                return

        new_password = simpledialog.askstring(
            "Profile password",
            "Enter a password to lock edits in Safe Mode. Leave blank to disable the lock.",
            show="*",
            parent=self,
        )
        if new_password is None:
            return
        self.profile_password = hash_profile_password(new_password.strip())
        self.save_current_rules()
        self.refresh_profile_status()
        if self.current_profile == "child_protection" and self.profile_password:
            show_app_dialog(self, "Safe Mode locked", "The profile is now locked. An admin password is required before editing the block list.", "success")

    def handle_profile_change(self, _event=None):
        selected = PROFILE_VALUES.get(self.profile_var.get())
        if selected is None:
            self.profile_var.set(PROFILE_LABELS[self.current_profile])
            return
        if selected == self.current_profile:
            return

        if self.profile_password and selected != self.current_profile:
            password = simpledialog.askstring(
                "Admin password",
                f"Enter the admin password to change from {PROFILE_LABELS[self.current_profile]} to {PROFILE_LABELS[selected]}:",
                show="*",
                parent=self,
            )
            if not verify_profile_password(password, self.profile_password):
                self.profile_var.set(PROFILE_LABELS[self.current_profile])
                show_app_dialog(self, "Access blocked", "Only the admin password can change the protection profile.", "warning")
                return

        if selected == "child_protection" and not self.profile_password:
            answer = simpledialog.askstring(
                "Set Safe Mode password",
                "Enter a password to lock changes in Safe Mode. Leave blank to stay in Admin mode.",
                show="*",
                parent=self,
            )
            if answer is None or not answer.strip():
                self.profile_var.set(PROFILE_LABELS[self.current_profile])
                return
            self.profile_password = hash_profile_password(answer.strip())

        self.current_profile = selected
        self.save_current_rules()
        self.refresh_profile_status()

    def require_profile_password(self, action_name):
        if self.current_profile != "child_protection" or not self.profile_password:
            return True
        password = simpledialog.askstring(
            "Password required",
            f"Enter the admin password to {action_name} in Safe Mode:",
            show="*",
            parent=self,
        )
        if verify_profile_password(password, self.profile_password):
            return True
        show_app_dialog(self, "Access blocked", "This profile is locked. Only the admin password can edit website protection or the block list.", "warning")
        return False

    def refresh_listbox(self):
        self.listbox.delete(0, tk.END)
        for item in self.blocklist:
            self.listbox.insert(tk.END, item)
        self.count_label.config(text=f"{len(self.blocklist)} domain{'s' if len(self.blocklist) != 1 else ''}")
        if hasattr(self, "domain_stat_label"):
            self.domain_stat_label.config(text=str(len(self.blocklist)))
        if hasattr(self, "rule_stat_label"):
            self.rule_stat_label.config(text=str(sum(len(words) for words in self.categories.values())))

    def refresh_hosts_status(self):
        status, state = inspect_hosts_rules()
        message, _status_label = protection_status_for_display(status, state)
        self.hosts_status_label.config(text=message, style="Muted.TLabel" if state == "success" else "App.TLabel")

    def parse_domains(self, value, limit=None):
        domains = []
        seen = set(self.blocklist)
        for item in re.split(r"[\s,;]+", value):
            domain = normalize_domain(item)
            if is_valid_domain(domain) and not is_protected_domain(domain) and domain not in seen:
                domains.append(domain)
                seen.add(domain)
                if limit and len(domains) >= limit:
                    break
        return domains

    def extract_domains(self, text, limit=None):
        candidates = []
        for line in text.splitlines():
            content = line.split("#", 1)[0].strip()
            if content:
                candidates.extend(re.split(r"[\s,;]+", content))
        return self.parse_domains(" ".join(candidates), limit=limit)

    def check_url(self):
        site = self.check_var.get().strip()
        if not site:
            show_app_dialog(self, "Missing website", "Please enter a website or domain.", "warning")
            return

        result = analyze_site(site, self.blocklist, self.categories, load_dns_blocklist())
        record_activity(
            "manual_check",
            result["domain"],
            f"{result['level']} ({result['score']}/100); {result['recommendation']}",
        )
        if not result["domain"]:
            show_app_dialog(self, "Invalid website", "Enter a valid website or domain.", "warning")
            return

        reasons = []
        if result["dns_feed_blocked"]:
            reasons.append("listed in the Avantis DNS blocklist feed")
        elif result["blocked"]:
            reasons.append("already blocked or a subdomain of a blocked site")
        if result["categories"]:
            reasons.append(f"categories: {', '.join(result['categories'])}")
        if result["keywords"]:
            reasons.append(f"matched words: {', '.join(result['keywords'])}")
        if result["signals"]:
            reasons.append("Additional local risk indicators:\n- " + "\n- ".join(result["signals"]))
        reasons.append(f"recommendation: {result['recommendation']}")

        if reasons:
            if result["blocked"] or result["level"] == "High risk":
                dialog_kind = "risk"
            elif result["level"] == "Review":
                dialog_kind = "warning"
            else:
                dialog_kind = "success"
            show_app_dialog(
                self,
                f"{result['level']} ({result['score']}/100)",
                f"{result['domain']}\n\n" + "\n".join(reasons),
                dialog_kind,
            )
        else:
            show_app_dialog(self, f"{result['level']} ({result['score']}/100)", f"{result['domain']}\n\n{result['recommendation']}", "success")

    def open_safety_center(self):
        dialog = tk.Toplevel(self)
        dialog.title("Safety Center")
        dialog.geometry("780x560")
        dialog.minsize(650, 460)
        dialog.configure(bg="#eaf4f7")
        dialog.transient(self)
        if hasattr(self, "window_icon"):
            dialog.iconphoto(True, self.window_icon)

        header = tk.Frame(dialog, bg="#0b6478", padx=22, pady=16)
        header.pack(fill="x")
        tk.Label(header, text="Safety Center", bg="#0b6478", fg="#ffffff", font=("Segoe UI", 17, "bold")).pack(anchor="w")
        tk.Label(header, text="Local manual-check activity", bg="#0b6478", fg="#e2f5f7", font=("Segoe UI", 10)).pack(anchor="w", pady=(3, 0))

        body = ttk.Frame(dialog, style="App.TFrame", padding=16)
        body.pack(fill="both", expand=True)
        ttk.Label(body, text="Only manual checks are recorded. Successful browsing is not logged.", style="Muted.TLabel", wraplength=700).pack(anchor="w", pady=(0, 12))

        settings = ttk.Frame(body, style="Panel.TFrame", padding=10)
        settings.pack(fill="x", pady=(0, 10))
        ttk.Label(settings, text="Keep records for", style="Panel.TLabel").pack(side="left")
        activity = load_activity()
        selected_retention = next((label for label, hours in RETENTION_OPTIONS.items() if hours == activity["retention_hours"]), "24 hours")
        retention_var = tk.StringVar(value=selected_retention)
        retention_menu = ttk.Combobox(settings, textvariable=retention_var, values=list(RETENTION_OPTIONS), state="readonly", width=14)
        retention_menu.pack(side="left", padx=(8, 0))

        list_frame = ttk.Frame(body, style="Panel.TFrame", padding=10)
        list_frame.pack(fill="both", expand=True)
        activity_scroll = ttk.Scrollbar(list_frame, orient="vertical")
        activity_scroll.pack(side="right", fill="y")
        activity_list = tk.Listbox(list_frame, font=("Consolas", 9), bg="#f8fafc", fg="#16324f", selectbackground="#087f92", selectforeground="#ffffff", relief="flat", borderwidth=0, yscrollcommand=activity_scroll.set)
        activity_list.pack(side="left", fill="both", expand=True)
        activity_scroll.config(command=activity_list.yview)
        count_label = ttk.Label(body, text="", style="Muted.TLabel")
        count_label.pack(anchor="w", pady=(8, 0))

        def refresh_activity(*_):
            current = load_activity()
            activity_list.delete(0, tk.END)
            for event in reversed(current["events"]):
                timestamp = datetime.fromisoformat(event["timestamp"]).astimezone().strftime("%Y-%m-%d %H:%M")
                activity_list.insert(tk.END, f"{timestamp}  |  {event['type']}  |  {event['domain']}  |  {event['detail']}")
            count_label.config(text=f"{len(current['events'])} record{'s' if len(current['events']) != 1 else ''} stored locally")

        def update_retention(_event=None):
            current = load_activity()
            save_activity(current["events"], RETENTION_OPTIONS[retention_var.get()])
            refresh_activity()

        def clear_activity():
            current = load_activity()
            save_activity([], current["retention_hours"])
            refresh_activity()

        retention_menu.bind("<<ComboboxSelected>>", update_retention)
        refresh_activity()
        actions = ttk.Frame(body, style="App.TFrame")
        actions.pack(fill="x", pady=(10, 0))
        ttk.Button(actions, text="Clear Activity", command=clear_activity, style="Secondary.TButton").pack(side="left")
        ttk.Button(actions, text="Close", command=dialog.destroy, style="Secondary.TButton").pack(side="right")

        dialog.protocol("WM_DELETE_WINDOW", dialog.destroy)
        dialog.update_idletasks()
        self.update_idletasks()
        position_x = self.winfo_rootx() + (self.winfo_width() - dialog.winfo_width()) // 2
        position_y = self.winfo_rooty() + (self.winfo_height() - dialog.winfo_height()) // 2
        dialog.geometry(f"+{max(position_x, 0)}+{max(position_y, 0)}")
        dialog.grab_set()
        self.wait_window(dialog)

    def add_url(self):
        if not self.require_profile_password("add domains"):
            return
        site = self.add_var.get().strip()
        if not site:
            show_app_dialog(self, "Missing website", "Type a website first.", "warning")
            return

        domains = self.parse_domains(site)
        if not domains:
            show_app_dialog(self, "Invalid domain", "Enter a valid domain or URL.", "warning")
            return

        self.blocklist.extend(domains)
        self.save_current_rules()
        self.refresh_listbox()
        self.add_var.set("")
        show_app_dialog(self, "Domain added", f"{len(domains)} domain{'s' if len(domains) != 1 else ''} added to the block list.", "success")

    def open_site_dictionary(self):
        dialog = tk.Toplevel(self)
        dialog.title("Site and word dictionary")
        dialog.geometry("820x620")
        dialog.minsize(680, 500)
        dialog.configure(bg="#eaf4f7")
        dialog.transient(self)
        if hasattr(self, "window_icon"):
            dialog.iconphoto(True, self.window_icon)

        header = tk.Frame(dialog, bg="#0b6478", padx=22, pady=16)
        header.pack(fill="x")
        tk.Label(header, text="Site and word dictionary", bg="#0b6478", fg="#ffffff", font=("Segoe UI", 17, "bold")).pack(anchor="w")
        tk.Label(header, text="Search blocked domains, detection words, and subdomain prefixes", bg="#0b6478", fg="#e2f5f7", font=("Segoe UI", 10)).pack(anchor="w", pady=(3, 0))

        body = ttk.Frame(dialog, style="App.TFrame", padding=16)
        body.pack(fill="both", expand=True)
        search_var = tk.StringVar()
        ttk.Label(body, text="Search the dictionary", style="Panel.TLabel", font=("Segoe UI", 10, "bold")).pack(anchor="w")
        search_entry = ttk.Entry(body, textvariable=search_var)
        search_entry.pack(fill="x", pady=(4, 10))

        summary = ttk.Label(body, text="", style="Muted.TLabel")
        summary.pack(anchor="w", pady=(0, 8))

        table_frame = ttk.Frame(body, style="Panel.TFrame", padding=8)
        table_frame.pack(fill="both", expand=True)
        table_scroll = ttk.Scrollbar(table_frame, orient="vertical")
        table_scroll.pack(side="right", fill="y")
        columns = ("type", "entry", "category", "meaning")
        table = ttk.Treeview(table_frame, columns=columns, show="headings", yscrollcommand=table_scroll.set)
        table.heading("type", text="Type")
        table.heading("entry", text="Domain or word")
        table.heading("category", text="Category")
        table.heading("meaning", text="Meaning")
        table.column("type", width=105, anchor="w")
        table.column("entry", width=220, anchor="w")
        table.column("category", width=130, anchor="w")
        table.column("meaning", width=310, anchor="w")
        table.pack(side="left", fill="both", expand=True)
        table_scroll.config(command=table.yview)

        entries = []
        for domain in sorted(set(self.blocklist)):
            categories = detect_categories(domain, self.categories)
            category_text = ", ".join(categories) if categories else "custom block"
            meaning = "Blocked domain and its configured subdomains."
            entries.append(("Domain", domain, category_text, meaning))
        for category, keywords in sorted(self.categories.items()):
            meaning = CATEGORY_DESCRIPTIONS.get(category, "Custom detection category.")
            for keyword in sorted(set(keywords)):
                entries.append(("Word", keyword, category, meaning))
        for prefix in sorted(set(self.subdomain_prefixes)):
            entries.append(("Prefix", prefix, "hostname", f"Common subdomain prefix such as {prefix}.example.com."))

        def refresh_dictionary(*_):
            query = search_var.get().strip().lower()
            table.delete(*table.get_children())
            visible = [entry for entry in entries if not query or any(query in str(value).lower() for value in entry)]
            for entry in visible:
                table.insert("", tk.END, values=entry)
            summary.config(text=f"{len(visible)} of {len(entries)} entries")

        search_var.trace_add("write", refresh_dictionary)
        refresh_dictionary()
        ttk.Button(body, text="Close", command=dialog.destroy, style="Secondary.TButton").pack(anchor="e", pady=(10, 0))

        dialog.protocol("WM_DELETE_WINDOW", dialog.destroy)
        dialog.update_idletasks()
        self.update_idletasks()
        position_x = self.winfo_rootx() + (self.winfo_width() - dialog.winfo_width()) // 2
        position_y = self.winfo_rooty() + (self.winfo_height() - dialog.winfo_height()) // 2
        dialog.geometry(f"+{max(position_x, 0)}+{max(position_y, 0)}")
        dialog.grab_set()
        search_entry.focus_set()
        self.wait_window(dialog)

    def open_rule_manager(self):
        if not self.require_profile_password("edit the rule library"):
            return
        dialog = tk.Toplevel(self)
        dialog.title("Manage search rules")
        dialog.geometry("680x560")
        dialog.minsize(580, 460)
        dialog.configure(bg="#eaf4f7")
        dialog.transient(self)
        if hasattr(self, "window_icon"):
            dialog.iconphoto(True, self.window_icon)

        header = tk.Frame(dialog, bg="#0b6478", padx=22, pady=16)
        header.pack(fill="x")
        tk.Label(header, text="Rule library", bg="#0b6478", fg="#ffffff", font=("Segoe UI", 17, "bold")).pack(anchor="w")
        tk.Label(header, text="Search, add, and refine the words Avantis FireWall detects", bg="#0b6478", fg="#e2f5f7", font=("Segoe UI", 10)).pack(anchor="w", pady=(3, 0))

        body = ttk.Frame(dialog, style="App.TFrame", padding=16)
        body.pack(fill="both", expand=True)
        search_var = tk.StringVar()
        ttk.Label(body, text="Search rules", style="Panel.TLabel", font=("Segoe UI", 10, "bold")).pack(anchor="w")
        search_entry = ttk.Entry(body, textvariable=search_var)
        search_entry.pack(fill="x", pady=(4, 12))

        list_frame = ttk.Frame(body, style="Panel.TFrame", padding=10)
        list_frame.pack(fill="both", expand=True)
        rule_scroll = ttk.Scrollbar(list_frame, orient="vertical")
        rule_scroll.pack(side="right", fill="y")
        rule_list = tk.Listbox(list_frame, selectmode=tk.EXTENDED, font=("Segoe UI", 10), bg="#f8fafc", fg="#16445a", selectbackground="#087f92", selectforeground="#ffffff", relief="flat", borderwidth=0, yscrollcommand=rule_scroll.set)
        rule_list.pack(side="left", fill="both", expand=True)
        rule_scroll.config(command=rule_list.yview)
        visible_rules = []

        def refresh_rules(*_):
            query = search_var.get().strip().lower()
            visible_rules.clear()
            rule_list.delete(0, tk.END)
            for category, keywords in sorted(self.categories.items()):
                for keyword in sorted(set(keywords)):
                    if not query or query in category.lower() or query in keyword.lower():
                        visible_rules.append((category, keyword))
                        rule_list.insert(tk.END, f"{category}  /  {keyword}")

        search_var.trace_add("write", refresh_rules)
        refresh_rules()

        editor = ttk.Frame(body, style="Panel.TFrame", padding=(0, 12, 0, 0))
        editor.pack(fill="x")
        category_var = tk.StringVar(value="gambling")
        keyword_var = tk.StringVar()
        ttk.Label(editor, text="Category", style="Panel.TLabel").grid(row=0, column=0, sticky="w")
        ttk.Label(editor, text="Keyword or phrase", style="Panel.TLabel").grid(row=0, column=1, sticky="w", padx=(12, 0))
        category_entry = ttk.Entry(editor, textvariable=category_var)
        category_entry.grid(row=1, column=0, sticky="ew", pady=(4, 0))
        keyword_entry = ttk.Entry(editor, textvariable=keyword_var)
        keyword_entry.grid(row=1, column=1, sticky="ew", padx=(12, 0), pady=(4, 0))
        editor.columnconfigure(0, weight=1)
        editor.columnconfigure(1, weight=2)

        actions = ttk.Frame(body, style="App.TFrame")
        actions.pack(fill="x", pady=(12, 0))

        def add_rule():
            if not self.require_profile_password("add a rule"):
                return
            category = category_var.get().strip().lower()
            keyword = keyword_var.get().strip().lower()
            if not category or not keyword:
                show_app_dialog(dialog, "Incomplete rule", "Enter both a category and a keyword.", "warning")
                return
            self.categories.setdefault(category, [])
            if keyword not in self.categories[category]:
                self.categories[category].append(keyword)
                self.save_current_rules()
            keyword_var.set("")
            refresh_rules()

        def remove_rules():
            if not self.require_profile_password("remove rules"):
                return
            selected = rule_list.curselection()
            if not selected:
                show_app_dialog(dialog, "No rules selected", "Select one or more rules to remove.", "warning")
                return
            for index in reversed(selected):
                category, keyword = visible_rules[index]
                self.categories[category] = [item for item in self.categories[category] if item != keyword]
                if not self.categories[category]:
                    del self.categories[category]
            self.save_current_rules()
            refresh_rules()

        ttk.Button(actions, text="Add Rule", command=add_rule, style="Accent.TButton").pack(side="left")
        ttk.Button(actions, text="Remove Selected", command=remove_rules, style="Secondary.TButton").pack(side="left", padx=(8, 0))
        ttk.Button(actions, text="Close", command=dialog.destroy, style="Secondary.TButton").pack(side="right")

        dialog.protocol("WM_DELETE_WINDOW", dialog.destroy)
        dialog.update_idletasks()
        self.update_idletasks()
        position_x = self.winfo_rootx() + (self.winfo_width() - dialog.winfo_width()) // 2
        position_y = self.winfo_rooty() + (self.winfo_height() - dialog.winfo_height()) // 2
        dialog.geometry(f"+{max(position_x, 0)}+{max(position_y, 0)}")
        dialog.grab_set()
        search_entry.focus_set()
        self.wait_window(dialog)

    def add_multiple(self):
        if not self.require_profile_password("import domains"):
            return
        value = self.bulk_text.get("1.0", tk.END).strip()
        if not value:
            show_app_dialog(self, "Nothing to import", "Paste at least one domain first.", "warning")
            return

        domains = self.parse_domains(value)
        if not domains:
            show_app_dialog(self, "Nothing to import", "No valid new domains were found.", "warning")
            return

        self.blocklist.extend(domains)
        self.save_current_rules()
        self.refresh_listbox()
        self.bulk_text.delete("1.0", tk.END)
        show_app_dialog(self, "Bulk import complete", f"{len(domains)} domain{'s' if len(domains) != 1 else ''} added to the block list.", "success")

    def import_file(self):
        if not self.require_profile_password("import domains from a file"):
            return
        file_path = filedialog.askopenfilename(
            parent=self,
            title="Import blocked sites",
            filetypes=[
                ("Supported files", "*.txt *.csv *.json *.docx"),
                ("Text files", "*.txt *.csv"),
                ("JSON files", "*.json"),
                ("Word documents", "*.docx"),
                ("All files", "*.*"),
            ],
        )
        if not file_path:
            return

        try:
            path = Path(file_path)
            if path.suffix.lower() == ".docx":
                document = Document(path)
                text_parts = [paragraph.text for paragraph in document.paragraphs]
                text_parts.extend(cell.text for table in document.tables for row in table.rows for cell in row.cells)
                text = "\n".join(text_parts)
            elif path.suffix.lower() == ".json":
                with path.open("r", encoding="utf-8") as file:
                    text = json.dumps(json.load(file))
            else:
                text = path.read_text(encoding="utf-8-sig", errors="ignore")

            domains = self.extract_domains(text)
            if not domains:
                show_app_dialog(self, "No sites found", "The selected file did not contain recognizable domain names.", "warning")
                return

            self.blocklist.extend(domains)
            self.save_current_rules()
            self.refresh_listbox()
            show_app_dialog(self, "File import complete", f"{len(domains)} new domain{'s' if len(domains) != 1 else ''} imported from:\n{path.name}", "success")
        except (OSError, UnicodeError, json.JSONDecodeError, ValueError) as error:
            show_app_dialog(self, "Import failed", str(error), "error")

    def remove_selected(self):
        if not self.require_profile_password("remove domains"):
            return
        selected = self.listbox.curselection()
        if not selected:
            show_app_dialog(self, "No domains selected", "Select one or more domains first.", "warning")
            return

        removed = [self.blocklist[index] for index in selected]
        for index in reversed(selected):
            self.blocklist.pop(index)
        self.save_current_rules()
        self.refresh_listbox()
        show_app_dialog(self, "Domains removed", f"{len(removed)} domain{'s' if len(removed) != 1 else ''} removed.", "success")

    def clear_all(self):
        if not self.require_profile_password("clear the blocked list"):
            return
        if show_app_dialog(self, "Clear blocked domains", "Remove all blocked domains?", "warning", confirm=True):
            self.blocklist.clear()
            self.save_current_rules()
            self.refresh_listbox()

    def apply_blocking(self):
        if len(self.blocklist) > HOSTS_DOMAIN_LIMIT:
            show_app_dialog(self, "Block list is too large", f"This list contains {len(self.blocklist)} domains. Remove some domains to meet the {HOSTS_DOMAIN_LIMIT}-domain limit before applying protection.", "warning")
            return

        if not is_admin():
            try:
                result = run_as_admin("apply")
                self.refresh_hosts_status()
                message = "Website protection is active on this PC."
                if not result["flushed"]:
                    message += "\n\nWindows could not refresh website access yet. Restart your browser or try again."
                show_app_dialog(self, "Protection applied", message, "success" if result["flushed"] else "warning")
            except Exception as error:
                show_app_dialog(self, "Could not apply protection", "Protection could not be updated. Check Administrator permissions and try again.", "error")
            return

        try:
            write_hosts_file(self.blocklist, self.subdomain_prefixes)
            self.refresh_hosts_status()
            flushed, flush_error = flush_dns_cache()
            message = "Website protection is active on this PC."
            if not flushed:
                message += "\n\nWindows could not refresh website access yet. Restart your browser or try again."
            show_app_dialog(self, "Protection applied", message, "success" if flushed else "warning")
        except Exception as e:
            self.refresh_hosts_status()
            show_app_dialog(self, "Could not apply protection", "Protection could not be updated. Check Administrator permissions and try again.", "error")

    def remove_blocking(self):
        if not show_app_dialog(self, "Remove website protection?", "This removes Avantis website restrictions from this PC. Domains covered by those rules may become accessible again. Administrator approval is required.", "warning", confirm=True):
            return

        if not is_admin():
            try:
                result = run_as_admin("remove")
                self.refresh_hosts_status()
                message = "Avantis website protection was removed from this PC." if result["removed"] else "There were no Avantis website restrictions to remove."
                if not result["flushed"]:
                    message += "\n\nWindows could not refresh website access yet. Restart your browser or try again."
                show_app_dialog(self, "Protection removed", message, "success" if result["flushed"] else "warning")
            except Exception as error:
                show_app_dialog(self, "Could not remove protection", "Protection could not be updated. Check Administrator permissions and try again.", "error")
            return

        try:
            removed = remove_hosts_rules()
            self.refresh_hosts_status()
            flushed, flush_error = flush_dns_cache()
            message = "Avantis website protection was removed from this PC." if removed else "There were no Avantis website restrictions to remove."
            if not flushed:
                message += f"\n\nWindows could not refresh website access yet: {flush_error}"
            show_app_dialog(self, "Protection removed" if removed else "Protection status checked", message, "success" if flushed else "warning")
        except Exception as e:
            self.refresh_hosts_status()
            show_app_dialog(self, "Could not remove protection", "Protection could not be updated. Check Administrator permissions and try again.", "error")

class DesktopApi:
    def __init__(self):
        config = load_config()
        self.blocklist = [normalize_domain(item) for item in config["blocked_domains"] if str(item).strip()]
        self.categories = config["categories"]
        self.subdomain_prefixes = config["subdomain_prefixes"]
        self.current_profile = config.get("protection_profile", "default")
        self.profile_password = str(config.get("profile_password", "") or "")
        self.window = None

    def attach_window(self, window):
        self.window = window

    def _save_rules(self, action="Policy settings updated"):
        previous_policy = policy_snapshot(load_config())
        current_policy = {
            "blocked_domains": self.blocklist,
            "categories": self.categories,
            "subdomain_prefixes": self.subdomain_prefixes,
        }
        save_rules(self.blocklist, self.categories, self.subdomain_prefixes, self.current_profile, self.profile_password)
        if previous_policy != current_policy:
            archive_policy({**previous_policy}, action)
        record_admin_audit(action, f"{len(self.blocklist)} blocked domains; {sum(len(words) for words in self.categories.values())} detection rules.")

    def _authorize(self, password, action):
        if self.current_profile == "child_protection" and self.profile_password and not verify_profile_password(password, self.profile_password):
            raise PermissionError(f"The admin password is required to {action} in Safe Mode.")

    def _parse_domains(self, value):
        domains = []
        seen = set(self.blocklist)
        for item in re.split(r"[\s,;]+", value):
            domain = normalize_domain(item)
            if is_valid_domain(domain) and not is_protected_domain(domain) and domain not in seen:
                domains.append(domain)
                seen.add(domain)
        return domains

    def get_policy_history(self):
        summaries = []
        for entry in reversed(load_policy_history()):
            policy = entry.get("policy", {})
            summaries.append({
                "id": entry.get("id", ""),
                "timestamp": entry.get("timestamp", ""),
                "action": entry.get("action", "Policy update"),
                "domain_count": len(policy.get("blocked_domains", [])),
                "rule_count": sum(len(words) for words in policy.get("categories", {}).values() if isinstance(words, list)),
            })
        return summaries

    def _validate_policy_document(self, document):
        policy = document.get("policy", document) if isinstance(document, dict) else None
        if not isinstance(policy, dict):
            raise ValueError("The selected file is not an Avantis policy file.")
        if isinstance(document, dict) and document.get("schema_version", 1) != 1:
            raise ValueError("This policy file uses an unsupported schema version.")
        domains = policy.get("blocked_domains")
        categories = policy.get("categories")
        prefixes = policy.get("subdomain_prefixes")
        if not isinstance(domains, list) or not isinstance(categories, dict) or not isinstance(prefixes, list):
            raise ValueError("The policy must contain domain, category, and subdomain-prefix lists.")

        clean_domains = []
        seen_domains = set()
        for value in domains:
            if not isinstance(value, str):
                raise ValueError("Every blocked-domain entry must be text.")
            domain = normalize_domain(value)
            if not is_valid_domain(domain):
                raise ValueError(f"Invalid domain in policy: {value}")
            if not is_protected_domain(domain) and domain not in seen_domains:
                clean_domains.append(domain)
                seen_domains.add(domain)

        clean_categories = {}
        for name, keywords in categories.items():
            category = str(name).strip().lower()
            if not category or not isinstance(keywords, list):
                raise ValueError("Every policy category must have a name and a list of keywords.")
            if any(not isinstance(keyword, str) for keyword in keywords):
                raise ValueError(f"Every detection rule in {category} must be text.")
            clean_categories[category] = list(dict.fromkeys(
                keyword.strip().lower() for keyword in keywords if keyword.strip()
            ))

        clean_prefixes = []
        for value in prefixes:
            if not isinstance(value, str):
                raise ValueError("Every subdomain prefix must be text.")
            prefix = value.strip().lower()
            if not re.fullmatch(r"[a-z0-9-]{1,63}", prefix):
                raise ValueError(f"Invalid subdomain prefix in policy: {value}")
            if prefix not in clean_prefixes:
                clean_prefixes.append(prefix)
        return {
            "blocked_domains": clean_domains,
            "categories": clean_categories,
            "subdomain_prefixes": clean_prefixes,
        }

    def import_policy(self, password=""):
        self._authorize(password, "import a policy")
        import webview

        selection = self.window.create_file_dialog(
            webview.OPEN_DIALOG,
            allow_multiple=False,
            file_types=("Avantis policy (*.json)",),
        )
        if not selection:
            return {"cancelled": True, "state": self.get_state()}
        selected_path = selection[0] if isinstance(selection, (list, tuple)) else selection
        try:
            document = json.loads(Path(selected_path).read_text(encoding="utf-8-sig"))
        except (OSError, UnicodeError, json.JSONDecodeError) as error:
            raise ValueError(f"Could not read the selected policy file: {error}") from error
        policy = self._validate_policy_document(document)
        self.blocklist = policy["blocked_domains"]
        self.categories = policy["categories"]
        self.subdomain_prefixes = policy["subdomain_prefixes"]
        self._save_rules("Policy imported")
        return {"cancelled": False, "state": self.get_state()}

    def restore_policy_version(self, snapshot_id, password=""):
        self._authorize(password, "restore a previous policy")
        snapshot = next((
            entry for entry in load_policy_history()
            if entry.get("id") == snapshot_id
        ), None)
        if not snapshot:
            raise ValueError("That policy version is no longer available.")
        policy = self._validate_policy_document(snapshot.get("policy"))
        self.blocklist = policy["blocked_domains"]
        self.categories = policy["categories"]
        self.subdomain_prefixes = policy["subdomain_prefixes"]
        self._save_rules("Previous policy restored")
        return self.get_state()

    def get_state(self):
        print("Avantis get_state request received.", flush=True)
        state = {
            "blocked_domains": self.blocklist,
            "categories": self.categories,
            "subdomain_prefixes": self.subdomain_prefixes,
            "protection_profile": self.current_profile,
            "has_profile_password": bool(self.profile_password),
            "hosts_status": inspect_hosts_rules(),
            "activity": load_activity(),
            "policy_history": self.get_policy_history(),
            "admin_audit": load_admin_audit(),
        }
        print("Avantis get_state request completed.", flush=True)
        return state

    def analyze_site(self, site):
        result = analyze_site(site, self.blocklist, self.categories, load_dns_blocklist())
        record_activity("manual_check", result["domain"], f"{result['level']} ({result['score']}/100); {result['recommendation']}")
        return result

    def add_domains(self, value, password=""):
        self._authorize(password, "add domains")
        domains = self._parse_domains(value)
        if not domains:
            raise ValueError("No valid new domains were found.")
        self.blocklist.extend(domains)
        self._save_rules("Domains added")
        return {"added": domains, "state": self.get_state()}

    def import_file(self, password=""):
        self._authorize(password, "import domains")
        import webview

        selection = self.window.create_file_dialog(
            webview.OPEN_DIALOG,
            allow_multiple=False,
            file_types=("Supported files (*.txt;*.csv;*.json;*.docx)", "All files (*.*)"),
        )
        if not selection:
            return {"added": [], "state": self.get_state()}

        selected_path = selection[0] if isinstance(selection, (list, tuple)) else selection
        path = Path(selected_path)
        if path.suffix.lower() == ".docx":
            document = Document(path)
            text_parts = [paragraph.text for paragraph in document.paragraphs]
            text_parts.extend(cell.text for table in document.tables for row in table.rows for cell in row.cells)
            text = "\n".join(text_parts)
        elif path.suffix.lower() == ".json":
            with path.open("r", encoding="utf-8") as file:
                text = json.dumps(json.load(file))
        else:
            text = path.read_text(encoding="utf-8-sig", errors="ignore")

        domains = self._parse_domains(text)
        if not domains:
            raise ValueError("The selected file did not contain any new valid domains.")
        self.blocklist.extend(domains)
        self._save_rules("Domains imported")
        return {"added": domains, "source": path.name, "state": self.get_state()}

    def change_profile(self, profile, password="", new_password=""):
        if profile not in {"default", "child_protection"}:
            raise ValueError("Unknown protection profile.")
        if profile == self.current_profile:
            return self.get_state()
        if self.profile_password and not verify_profile_password(password, self.profile_password):
            raise PermissionError("The admin password is required to change the protection profile.")
        if profile == "child_protection" and not self.profile_password:
            if not new_password.strip():
                raise ValueError("Set a password before enabling Safe Mode.")
            self.profile_password = hash_profile_password(new_password.strip())
        self.current_profile = profile
        self._save_rules("Protection profile changed")
        return self.get_state()

    def admin_access(self, password=""):
        if self.current_profile == "child_protection" and self.profile_password and not verify_profile_password(password, self.profile_password):
            raise PermissionError("The admin password is incorrect.")
        return {
            "profile": self.current_profile,
            "has_profile_password": bool(self.profile_password),
            "blocked_count": len(self.blocklist),
        }

    def set_profile_password(self, current_password, new_password):
        if self.profile_password and not verify_profile_password(current_password, self.profile_password):
            raise PermissionError("The current admin password is incorrect.")
        self.profile_password = hash_profile_password(new_password.strip())
        self._save_rules("Admin password settings changed")
        return self.get_state()

    def restore_admin(self, password=""):
        if self.profile_password and not verify_profile_password(password, self.profile_password):
            raise PermissionError("The admin password is incorrect.")
        self.current_profile = "default"
        self._save_rules("Admin mode restored")
        return self.get_state()

    def remove_domains(self, domains, password=""):
        self._authorize(password, "remove domains")
        selected = set(domains)
        self.blocklist = [domain for domain in self.blocklist if domain not in selected]
        self._save_rules("Domains removed")
        return self.get_state()

    def clear_domains(self, password=""):
        self._authorize(password, "clear the blocked list")
        self.blocklist.clear()
        self._save_rules("Blocked domains cleared")
        return self.get_state()

    def add_rule(self, category, keyword, password=""):
        self._authorize(password, "edit the rule library")
        category = category.strip().lower()
        keyword = keyword.strip().lower()
        if not category or not keyword:
            raise ValueError("Enter both a category and a keyword.")
        self.categories.setdefault(category, [])
        if keyword not in self.categories[category]:
            self.categories[category].append(keyword)
            self._save_rules("Detection rule added")
        return self.get_state()

    def remove_rules(self, rules, password=""):
        self._authorize(password, "edit the rule library")
        for rule in rules:
            category = rule.get("category")
            keyword = rule.get("keyword")
            if category in self.categories:
                self.categories[category] = [item for item in self.categories[category] if item != keyword]
                if not self.categories[category]:
                    del self.categories[category]
        self._save_rules("Detection rules removed")
        return self.get_state()

    def get_dictionary(self):
        entries = []
        for domain in sorted(set(self.blocklist)):
            categories = detect_categories(domain, self.categories)
            entries.append({"type": "Domain", "entry": domain, "category": ", ".join(categories) if categories else "custom block", "meaning": "Blocked domain and its configured subdomains."})
        for category, keywords in sorted(self.categories.items()):
            meaning = CATEGORY_DESCRIPTIONS.get(category, "Custom detection category.")
            entries.extend({"type": "Word", "entry": keyword, "category": category, "meaning": meaning} for keyword in sorted(set(keywords)))
        entries.extend({"type": "Prefix", "entry": prefix, "category": "hostname", "meaning": f"Common subdomain prefix such as {prefix}.example.com."} for prefix in sorted(set(self.subdomain_prefixes)))
        return entries

    def update_retention(self, retention_hours):
        retention_hours = int(retention_hours)
        if retention_hours not in RETENTION_OPTIONS.values():
            raise ValueError("Choose a supported activity retention period.")
        activity = load_activity()
        save_activity(activity["events"], retention_hours)
        record_admin_audit("Activity retention changed", f"Retention set to {retention_hours} hours.")
        return load_activity()

    def clear_activity(self):
        activity = load_activity()
        save_activity([], activity["retention_hours"])
        record_admin_audit("Manual-check activity cleared", "Stored manual-check records were cleared.")
        return load_activity()

    def apply_hosts(self):
        if os.name == "nt" and not is_admin():
            return run_as_admin("apply")
        if len(self.blocklist) > HOSTS_DOMAIN_LIMIT:
            raise ValueError(f"Windows hosts mode is limited to {HOSTS_DOMAIN_LIMIT} domains. Use a managed DNS service for larger lists.")
        write_hosts_file(self.blocklist, self.subdomain_prefixes)
        flushed, error = flush_dns_cache()
        record_admin_audit("Website protection applied", f"{len(self.blocklist)} website rules applied.")
        return {"flushed": flushed, "error": error, "hosts_status": inspect_hosts_rules()}

    def remove_hosts(self):
        if os.name == "nt" and not is_admin():
            return run_as_admin("remove")
        removed = remove_hosts_rules()
        flushed, error = flush_dns_cache()
        record_admin_audit("Website protection removed", "Avantis website restrictions were removed." if removed else "No Avantis website restrictions were active.")
        return {"removed": removed, "flushed": flushed, "error": error, "hosts_status": inspect_hosts_rules()}

def main():
    if not show_first_run_agreement():
        return

    initialize_data_root()
    os.environ.setdefault("QTWEBENGINE_CHROMIUM_FLAGS", "--disable-gpu")
    import webview

    api = DesktopApi()
    interface_path = (RESOURCE_ROOT / "web" / "dist" / "index.html").resolve()
    if not interface_path.is_file():
        raise FileNotFoundError("The React interface is not built yet. Run npm.cmd install and npm.cmd run build from the web folder.")

    window = webview.create_window(
        "Avantis FireWall",
        str(interface_path),
        js_api=api,
        min_size=(860, 640),
        width=1180,
        height=820,
    )
    api.attach_window(window)
    print(f"Avantis loading interface: {interface_path}", flush=True)
    window.events.loaded += lambda: print("Avantis React document loaded.", flush=True)
    window.events._pywebviewready += lambda: print("Avantis Python bridge ready.", flush=True)
    webview.start(gui="qt", icon=str(APP_ICON_FILE) if APP_ICON_FILE.exists() else None)


def show_first_run_agreement():
    try:
        with AGREEMENT_FILE.open("r", encoding="utf-8") as f:
            accepted = json.load(f)
        if accepted.get("version") == AGREEMENT_VERSION and accepted.get("accepted") is True:
            return True
    except (OSError, json.JSONDecodeError, AttributeError):
        pass

    root = tk.Tk()
    root.title("Avantis FireWall - Terms and Privacy")
    root.geometry("760x600")
    root.minsize(640, 520)
    root.configure(background="#f4f9fb")
    if APP_ICON_FILE.exists():
        try:
            root.iconbitmap(str(APP_ICON_FILE))
        except tk.TclError:
            pass

    accepted = []
    agreed = tk.BooleanVar(value=False)
    outer = ttk.Frame(root, padding=(28, 24))
    outer.pack(fill="both", expand=True)
    ttk.Label(outer, text="AVANTIS FIREWALL", foreground="#0b8194", font=("Segoe UI", 10, "bold")).pack(anchor="w")
    ttk.Label(outer, text="Before you continue", foreground="#173a50", font=("Segoe UI", 22, "bold")).pack(anchor="w", pady=(8, 4))
    ttk.Label(outer, text="Please review and accept these terms and the privacy notice.", foreground="#617c8d", font=("Segoe UI", 10)).pack(anchor="w", pady=(0, 18))

    text_frame = ttk.Frame(outer)
    text_frame.pack(fill="both", expand=True)
    policy_text = tk.Text(text_frame, wrap="word", height=14, padx=16, pady=14, relief="solid", borderwidth=1, background="#ffffff", foreground="#294c61", font=("Segoe UI", 10), spacing1=2, spacing3=5)
    scrollbar = ttk.Scrollbar(text_frame, orient="vertical", command=policy_text.yview)
    policy_text.configure(yscrollcommand=scrollbar.set)
    policy_text.pack(side="left", fill="both", expand=True)
    scrollbar.pack(side="right", fill="y")
    policy_text.tag_configure("heading", font=("Segoe UI", 11, "bold"), foreground="#173f52", spacing1=12, spacing3=5)
    policy_text.insert("end", "TERMS OF USE\n", "heading")
    policy_text.insert("end", "Use Avantis only on a device you own or are authorized to administer. Protection changes can affect everyone using this PC. You are responsible for choosing appropriate rules and complying with applicable law and workplace or family policies.\n\n")
    policy_text.insert("end", "HOW PROTECTION WORKS\n", "heading")
    policy_text.insert("end", "Check a website runs local rule-based analysis and reports its findings; it does not change access. Add a domain to your block list and select Apply Protection to restrict access to it and its configured common subdomains. Rules apply across browsers and Windows accounts on this PC. A listed domain may be blocked even if it is legitimate or safe. Remove the domain and apply the updated list, or select Remove Protection, to restore access.\n\n")
    policy_text.insert("end", "ADMINISTRATOR ACCESS\n", "heading")
    policy_text.insert("end", "Avantis normally opens without administrator approval. Windows requests Administrator approval only when you choose to apply or remove website protection. Approve that prompt only if you are authorized to administer this PC. Applying or removing protection changes website access for every user of this PC. You can remove Avantis protection from the app at any time.\n\n")
    policy_text.insert("end", "SAFE MODE AND PASSWORD\n", "heading")
    policy_text.insert("end", "Safe Mode is an optional restricted profile that limits editing and hides administration pages. When you first switch to Safe Mode without a password, Avantis asks you to create one. Use the Safe Mode password control to set or change it. When set, the password is required to unlock administrator controls or return to Admin mode. Choose a unique password and keep it available: it is stored locally with this PC's settings, is not an online account password, and is not included in exported policy files. Removing the password removes the password gate.\n\n")
    policy_text.insert("end", "PRIVACY NOTICE\n", "heading")
    policy_text.insert("end", "Avantis does not require an account. Manual checks, protection settings, and administrative actions are kept on this device; normal browsing is not monitored. If you choose to download a public blocklist, your device connects to that list's provider. Avantis does not send your browsing history to that provider.\n\n")
    policy_text.insert("end", "Policy backup files are created only when you request them and do not include your Safe Mode password. Manual-check activity can be cleared in the app; administrative actions remain available locally for review.\n\n")
    policy_text.insert("end", "By selecting “I agree” and continuing, you confirm that you have read and accept these Terms of Use and acknowledge this Privacy Notice.")
    policy_text.configure(state="disabled")

    ttk.Checkbutton(outer, text="I agree to the Terms of Use and acknowledge the Privacy Notice.", variable=agreed).pack(anchor="w", pady=(16, 12))
    buttons = ttk.Frame(outer)
    buttons.pack(fill="x")
    continue_button = ttk.Button(buttons, text="Agree and continue", state="disabled")
    continue_button.pack(side="right")
    ttk.Button(buttons, text="Decline and exit", command=root.destroy).pack(side="right", padx=(0, 8))

    def update_continue(*_):
        continue_button.configure(state="normal" if agreed.get() else "disabled")

    def accept_terms():
        if not agreed.get():
            return
        try:
            AGREEMENT_FILE.parent.mkdir(parents=True, exist_ok=True)
            with AGREEMENT_FILE.open("w", encoding="utf-8") as f:
                json.dump({"version": AGREEMENT_VERSION, "accepted": True, "accepted_at": datetime.now(timezone.utc).isoformat()}, f, indent=2)
        except OSError as error:
            messagebox.showerror("Could not save your agreement", "Avantis could not save your acceptance. Please try again.", parent=root)
            return
        accepted.append(True)
        root.destroy()

    agreed.trace_add("write", update_continue)
    continue_button.configure(command=accept_terms)
    root.bind("<Return>", lambda _event: accept_terms() if agreed.get() else None)
    root.protocol("WM_DELETE_WINDOW", root.destroy)
    root.mainloop()
    return bool(accepted)


if __name__ == "__main__":
    if len(sys.argv) == 4 and sys.argv[1] == "--hosts-helper":
        sys.exit(run_hosts_helper(sys.argv[2], sys.argv[3]))
    main()
